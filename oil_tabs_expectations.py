"""
Oil tab: Expectations (market-implied, published-model and model-implied forward-looking indicators), their link to oil, and a grade table.
Executed by oil_main.py. The indicators come from silver_expectations.py (shared engine); only the textbook signs differ for oil.
"""
import math
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

_XC = [c for c in EXP_KEYS if c in F_full.columns]
_ICON = {"Market-implied": "🟦", "Published model": "🟨", "Model-implied": "🟪"}
_STXT = {1: "Positive", -1: "Negative", 0: "Mixed"}
_ACTUAL = {"ex_ff12": ("fed_funds", "Actual Fed funds (today)"), "ex_y10": ("dgs10", "Actual 10Y yield (today)"),
           "ex_growth": ("gdp_yoy", "Actual GDP growth (latest)"), "ex_unrate": ("unrate_lvl", "Actual unemployment (latest)"),
           "ex_ind": ("indpro_yoy", "Actual industrial production YoY (latest)"), "ex_infl": ("cpi_yoy", "Actual CPI inflation YoY (latest)")}


def _fx(c, v, delta=False):
    if pd.isna(v):
        return "-"
    f = EXP_INFO[c]["fmt"]
    if f == "usd":
        return f"{v * 100:+.1f} pts" if delta else f"{v:+.1%}"
    if f == "prob":
        return f"{v:+.0f} pts" if delta else f"{v:.0f}%"
    if f == "chg":
        return f"{v:+.2f} pts"
    return f"{v:+.2f} pts" if delta else f"{v:.2f}%"


def _grade(m, E, idx, h=EXP_H):
    """Score each model/market expectation against what happened h months later versus 'nothing changes'. All values are point-in-time, so grading is out-of-sample."""
    def g(k):
        return m[k] if k in m else pd.Series(np.nan, index=m.index)
    ff = g("fed_funds").fillna(g("dgs3m"))
    usd = g("dollar")
    spec = {"ex_ff12": (ff, ff.shift(-h)), "ex_y10": (g("dgs10"), g("dgs10").shift(-h)), "ex_growth": (g("gdp_yoy"), g("gdp_yoy").shift(-h)),
            "ex_unrate": (g("unrate_lvl"), g("unrate_lvl").shift(-h)), "ex_usd": (pd.Series(0.0, index=m.index), usd.shift(-h) / usd - 1),
            "ex_ind": (g("indpro_yoy"), g("indpro_yoy").shift(-h))}
    rows = []
    for k, (anchor, real) in spec.items():
        if k not in E.columns:
            continue
        d = pd.concat([E[k], real, anchor], axis=1, keys=["p", "r", "a"]).reindex(idx).dropna()
        if len(d) < 24:
            continue
        mae_m, mae_n = (d["p"] - d["r"]).abs().mean(), (d["a"] - d["r"]).abs().mean()
        rows.append(dict(k=k, n=len(d), mae_m=mae_m, mae_n=mae_n, skill=1 - mae_m / mae_n if mae_n > 0 else np.nan,
                         ic=rank_ic(d["p"] - d["a"], d["r"] - d["a"])))
    return pd.DataFrame(rows)


with T["xpt"]:
    st.subheader("📈 Expectations: what rates, inflation, growth, jobs and the dollar are priced / projected to do over the next 12 months")
    st.caption("Ten forward-looking indicators, all **point-in-time**. 🟦 **Market-implied** = from today's Treasury curve, nothing fitted. 🟨 **Published model** = Cleveland Fed. "
               "🟪 **Model-implied** = walk-forward ridge estimate (no free point-in-time survey history exists). Oil lens: expected rates and inflation matter through the dollar and real yields, "
               "expected growth and industrial output through demand.")
    if not _XC:
        st.error("No expectation indicator could be built (the Treasury / FRED series did not download). Click 'Refresh data' or add a FRED_API_KEY.")
    else:
        last = F_full[_XC].dropna(how="all").index[-1]
        st.markdown(f"**Latest read: {last:%b %Y}.** Change shown versus 3 months earlier.")
        for row_cols in (_XC[0:4], _XC[4:7], _XC[7:10]):
            for col_, c in zip(st.columns(len(row_cols)), row_cols):
                s_ = F_full[c].dropna()
                box = card(col_)
                if s_.empty:
                    box.markdown(f"**{META_ALL[c][0]}**")
                    box.caption("not available")
                    continue
                prev = F_full[c].shift(3).loc[s_.index[-1]]
                box.metric(META_ALL[c][0], _fx(c, s_.iloc[-1]), None if pd.isna(prev) else _fx(c, s_.iloc[-1] - prev, True), delta_color="off")
                box.caption(f"{_ICON[EXP_INFO[c]['method']]} {EXP_INFO[c]['method']}")

        def val(c):
            s_ = F_full[c].dropna() if c in F_full else pd.Series(dtype=float)
            return s_.iloc[-1] if len(s_) else np.nan

        lines = []
        if pd.notna(val("ex_ff_chg")):
            d_ff = val("ex_ff_chg")
            lines.append(f"**Fed:** the curve prices {'cuts' if d_ff <= -0.25 else ('hikes' if d_ff >= 0.25 else 'roughly no change')} ({d_ff:+.2f} pts over 12 months).")
        if pd.notna(val("ex_real")) and pd.notna(val("ex_infl")):
            lines.append(f"**Real policy rate in 12 months:** about {val('ex_real'):.2f}% (expected inflation {val('ex_infl'):.2f}%).")
        if pd.notna(val("ex_recess")):
            rp = val("ex_recess")
            lines.append(f"**Recession risk (yield-curve model):** {rp:.0f}% ({'low' if rp < 20 else ('moderate' if rp < 40 else 'elevated')}). For oil, a recession means demand loss.")
        if pd.notna(val("ex_ind")):
            lines.append(f"**Industrial demand (model-implied):** production growth about {val('ex_ind'):.1f}% in 12 months.")
        if pd.notna(val("ex_usd")):
            u_ = val("ex_usd")
            lines.append(f"**Dollar (model-implied):** {'stronger' if u_ >= 0.01 else ('weaker' if u_ <= -0.01 else 'about flat')} ({u_:+.1%}; a low-confidence kind of forecast).")
        if lines:
            st.markdown("\n".join(f"- {x}" for x in lines))

        sup = uns = neu = mix = 0
        for c in _XC:
            sg = OIL_EXP_SIGN[c]
            p_ = P_full[c].dropna().iloc[-1] if c in P_full and P_full[c].notna().any() else np.nan
            if sg == 0:
                mix += 1
            elif pd.isna(p_):
                continue
            elif (p_ >= 2 / 3 and sg > 0) or (p_ <= 1 / 3 and sg < 0):
                sup += 1
            elif (p_ >= 2 / 3 and sg < 0) or (p_ <= 1 / 3 and sg > 0):
                uns += 1
            else:
                neu += 1
        st.markdown(f"**Textbook read for oil:** {sup} expectation(s) point in oil's favor, {uns} against, {neu} in the middle of their range, {mix} without a clear sign. "
                    "'Textbook' = lower real rates, a weaker dollar, higher inflation expectations, stronger growth and industrial demand and lower recession risk help oil. "
                    "A checklist, not evidence: the table below shows what THIS sample says.")

        st.markdown("#### History")
        pk = st.selectbox("Indicator", _XC, format_func=lambda c: META_ALL[c][0], key="oil_xpt_pick")
        s_ = F_full[pk].loc[start:].dropna()
        if s_.empty:
            st.info("No history for this indicator after the chosen start date.")
        else:
            k_ = 100.0 if EXP_INFO[pk]["fmt"] == "usd" else 1.0
            f = go.Figure()
            f.add_scatter(x=s_.index, y=s_.values * k_, name=META_ALL[pk][0], line=dict(color="#1f5fbf", width=2.5))
            if pk in _ACTUAL and _ACTUAL[pk][0] in m:
                a_ = m[_ACTUAL[pk][0]].reindex(s_.index)
                f.add_scatter(x=a_.index, y=a_.values, name=_ACTUAL[pk][1], line=dict(color="#c0392b", dash="dot", width=1.5))
            o_ = m["oil"].reindex(s_.index)
            f.add_scatter(x=o_.index, y=o_.values, name="WTI, USD/bbl (right axis, log)", yaxis="y2", line=dict(color="#999", width=1.5))
            f.update_layout(title=META_ALL[pk][0], height=400, yaxis2=dict(overlaying="y", side="right", showgrid=False, type="log"),
                            legend=dict(orientation="h", y=-0.2))
            show_plot(st, f)
            st.caption(f"{_ICON[EXP_INFO[pk]['method']]} **{EXP_INFO[pk]['method']}.** {EXP_INFO[pk]['how']}")

        st.markdown(f"#### Do these expectations line up with oil's {h}-month outcome? ({target_name.lower()})")
        rows = []
        for c in _XC:
            x_ = F_full[c].reindex(fwd.index)
            xl, yl, xt, yt = x_.loc[learn_idx], fwd.loc[learn_idx], x_.loc[unseen_idx], fwd.loc[unseen_idx]
            ic_l, ic_t = rank_ic(xl, yl), rank_ic(xt, yt)
            dl = pd.concat([xl, yl], axis=1).dropna()
            t_l = nw_t(dl.iloc[:, 1].rank().values, dl.iloc[:, 0].rank().values, h) if len(dl) > 8 else np.nan
            same = bool(pd.notna(ic_l) and pd.notna(ic_t) and np.sign(ic_l) == np.sign(ic_t))
            status = "-" if (pd.isna(ic_l) or pd.isna(ic_t)) else ("❌ Flips on unseen data" if not same else ("✅ Consistent and significant" if abs(t_l) >= 2 else "🟡 Consistent but weak"))
            sg = OIL_EXP_SIGN[c]
            rows.append({"Indicator": META_ALL[c][0], "Method": f"{_ICON[EXP_INFO[c]['method']]} {EXP_INFO[c]['method']}", "Learn months": len(dl),
                         "Rank IC (learn)": "-" if pd.isna(ic_l) else f"{ic_l:+.2f}", f"Rank IC ({unseen_lbl})": "-" if pd.isna(ic_t) else f"{ic_t:+.2f}",
                         "HAC t (learn)": "-" if pd.isna(t_l) else f"{t_l:+.1f}", "Status": status, "Textbook sign": _STXT[sg],
                         "Matches textbook?": "n/a" if sg == 0 or pd.isna(ic_l) else ("✔" if np.sign(ic_l) == sg else "✘")})
        show_df(st, pd.DataFrame(rows))
        st.caption("Expectations built from the same curve (Fed path, real rate, 10Y, recession) are strongly correlated with each other and with the dollar / real-yield indicators, "
                   "so they are not independent votes. Ten indicators are tested here, so a lone ✅ can be luck.")

        st.markdown("#### Are the expectations any good? (what actually happened 12 months later)")
        gr = _grade(m, F_full[_XC], ev_r)
        if gr.empty:
            st.info("Not enough overlapping months to grade the expectations.")
        else:
            nm = {"ex_ff12": "Fed funds, 12M ahead", "ex_y10": "10Y yield, 12M ahead", "ex_growth": "GDP growth", "ex_unrate": "Unemployment rate",
                  "ex_usd": "Dollar 12M change", "ex_ind": "Industrial production growth"}
            show_df(st, pd.DataFrame({
                "Expectation": [f"{nm[k]} {_ICON[EXP_INFO[k]['method']]}" for k in gr["k"]], "Months graded": gr["n"],
                "Avg error, expectation": gr["mae_m"].map(lambda v: f"{v:.3f}"), "Avg error, 'nothing changes'": gr["mae_n"].map(lambda v: f"{v:.3f}"),
                "Skill vs 'nothing changes'": gr["skill"].map(pc), "IC of expected vs realised change": gr["ic"].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}")}))
            beat = int((gr["skill"] > 0).sum())
            (st.success if beat >= len(gr) * 0.6 else st.warning)(
                f"{beat} of {len(gr)} expectations had a smaller average error than 'nothing changes'. Windows overlap, so independent observations are roughly months ÷ 12. "
                + ("Where skill is negative, treat the indicator as a description of the model's view, not as a forecast." if beat < len(gr) else ""))
            st.caption("Recession probability and expected inflation are not graded here (they need official recession dates / are published series).")

# ====================================================================== Options & geopolitics (live snapshot)
import os  # noqa: E402
import yfinance as yf  # noqa: E402

OPT_LOG = os.path.join(_HERE0, "oil_options_log.csv")


@st.cache_data(show_spinner=False, ttl=3600)
def options_snapshot(symbol="USO"):
    """Live USO option-chain readings. A SINGLE snapshot: Yahoo has no option history, so this cannot be graded or used in the rules.
    atm_near / atm_far = at-the-money implied vol (%) for a ~1-2 month and a ~3-6 month expiry; term = far - near (negative = near-term stress);
    skew = IV of a put ~10% out of the money minus IV of a call ~10% out (positive = downside hedging, negative = upside / supply-scare demand);
    pc_vol / pc_oi = put-to-call ratios of volume and open interest in the near expiry."""
    t = yf.Ticker(symbol)
    spot = float(t.history(period="5d")["Close"].dropna().iloc[-1])
    exps = list(t.options)
    if not exps:
        raise ValueError("no option expiries returned")
    today = pd.Timestamp.today().normalize()
    dte = {e: (pd.Timestamp(e) - today).days for e in exps}
    near = next((e for e in exps if 14 <= dte[e] <= 60), exps[0])
    far = next((e for e in exps if 75 <= dte[e] <= 200), None)

    def iv_at(df, k):
        d = df[(df["impliedVolatility"] > 0.01) & (df["openInterest"].fillna(0) + df["volume"].fillna(0) > 0)]
        if d.empty:
            return np.nan
        return float(d.iloc[(d["strike"] - k).abs().argsort().iloc[0]]["impliedVolatility"]) * 100

    oc = t.option_chain(near)
    c, p = oc.calls, oc.puts
    out = dict(symbol=symbol, spot=spot, near=near, near_dte=dte[near], far=far, far_dte=dte.get(far),
               atm_near=np.nanmean([iv_at(c, spot), iv_at(p, spot)]), skew=iv_at(p, 0.9 * spot) - iv_at(c, 1.1 * spot),
               pc_vol=float(p["volume"].fillna(0).sum() / max(c["volume"].fillna(0).sum(), 1)),
               pc_oi=float(p["openInterest"].fillna(0).sum() / max(c["openInterest"].fillna(0).sum(), 1)), atm_far=np.nan)
    if far:
        oc2 = t.option_chain(far)
        out["atm_far"] = np.nanmean([iv_at(oc2.calls, spot), iv_at(oc2.puts, spot)])
    out["term"] = out["atm_far"] - out["atm_near"] if pd.notna(out["atm_far"]) else np.nan
    return out


def _opt_read(o):
    """Plain-English reading of a snapshot."""
    r = []
    if pd.notna(o["term"]):
        r.append("Implied vol is **higher in the near expiry than the far one** (inverted): the market prices a near-term shock, typical of disruption scares." if o["term"] < -1 else
                 "Implied vol rises with expiry (normal shape): no sign of an acute near-term scare." if o["term"] > 1 else "The implied-vol term structure is flat.")
    if pd.notna(o["skew"]):
        r.append("**Calls are richer than puts** (negative skew): demand for upside protection, the pattern seen when supply disruption is feared." if o["skew"] < -1 else
                 "**Puts are richer than calls**: investors are paying for downside protection (demand / growth worries)." if o["skew"] > 1 else "Put and call wings are priced about equally.")
    r.append(f"Put/call ratio {o['pc_vol']:.2f} by volume and {o['pc_oi']:.2f} by open interest (above 1 = more puts).")
    return r


with T["xpt"]:
    st.markdown("---")
    st.markdown("#### 🎯 Options market: what traders are paying for (geopolitics, demand shocks)")
    st.caption("Options prices are a direct, forward-looking read on fear. **Historical** part (used in the rules, Market structure pillar): OVX, the options-implied volatility of oil, and "
               "its gap to realized volatility (`Oil options: implied minus realized`). **Live** part below: a snapshot of USO options. Yahoo keeps NO option history, so the snapshot cannot be "
               "backtested or graded. You can save it daily with the button and, after a year or two, the log becomes a usable indicator.")
    ov = F_full["ovx"].dropna() if "ovx" in F_full else pd.Series(dtype=float)
    if len(ov):
        pov = P_full["ovx"].dropna().iloc[-1] if "ovx" in P_full and P_full["ovx"].notna().any() else np.nan
        a, b, c_ = st.columns(3)
        a.metric("OVX (implied vol of oil)", f"{ov.iloc[-1]:.1f}", None if pd.isna(pov) else f"{pov:.0%} percentile vs past", delta_color="off")
        vr = F_full["ovx_vrp"].dropna() if "ovx_vrp" in F_full else pd.Series(dtype=float)
        if len(vr):
            b.metric("Implied minus realized vol", f"{vr.iloc[-1]:+.1f} pts", help="Positive = options pricing more movement than recently happened: a fear premium.")
        gp = F_full["gpr"].dropna() if "gpr" in F_full else pd.Series(dtype=float)
        if len(gp):
            c_.metric("Geopolitical risk index", f"{gp.iloc[-1]:.0f}")
    try:
        snap = options_snapshot("USO")
        s1, s2, s3, s4 = st.columns(4)
        s1.metric(f"USO ATM implied vol ({snap['near_dte']}d)", f"{snap['atm_near']:.1f}%")
        s2.metric("Term (far minus near)", "-" if pd.isna(snap["term"]) else f"{snap['term']:+.1f} pts")
        s3.metric("Put skew minus call skew", "-" if pd.isna(snap["skew"]) else f"{snap['skew']:+.1f} pts")
        s4.metric("Put / call (volume)", f"{snap['pc_vol']:.2f}")
        st.markdown("\n".join(f"- {x}" for x in _opt_read(snap)))
        if st.button("💾 Save today's snapshot to the log", key="oil_opt_save"):
            row = pd.DataFrame([dict(date=pd.Timestamp.today().normalize(), spot=snap["spot"], atm_near=snap["atm_near"], atm_far=snap["atm_far"],
                                     skew=snap["skew"], pc_vol=snap["pc_vol"], pc_oi=snap["pc_oi"])])
            old = pd.read_csv(OPT_LOG, parse_dates=["date"]) if os.path.exists(OPT_LOG) else pd.DataFrame()
            pd.concat([old, row]).drop_duplicates("date", keep="last").to_csv(OPT_LOG, index=False)
            st.success("Saved.")
        if os.path.exists(OPT_LOG):
            lg = pd.read_csv(OPT_LOG)
            st.caption(f"Log: {len(lg)} snapshot(s) in `oil_options_log.csv`. Needs about 24 months of month-end rows before it can feed the rules.")
    except Exception as e:  # noqa: BLE001
        snap = None
        st.info(f"Live USO option data could not be loaded ({type(e).__name__}). The historical OVX-based indicators still work.")
    st.caption("Limits: USO options are options on an ETF (they carry roll and ETF effects), not on WTI futures. Yahoo implied vols can be stale for illiquid strikes. "
               "CME's WTI futures options (the real source of skew / risk-reversal history) are not free. Treat the snapshot as context, not a signal.")

# ====================================================================== Report (Word)
# Document-building helpers (same as the silver report), then one section function per tab.
import io
import re

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    _HAS_DOCX = True
except Exception:  # noqa: BLE001
    _HAS_DOCX = False

_RP_REPL = {"✅": "Yes", "❌": "No", "✔": "Yes", "✘": "No", "🟢": "(+)", "🔴": "(-)", "⚪": "(0)", "🟡": "(~)"}
_RP_EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F\u200d]")
_RP_IMG_OK = [True]
_RP_VOTE = {1: "favorable", 0: "neutral", -1: "unfavorable"}


def _rp_clean(x):
    s = "" if x is None else str(x)
    for a, b in _RP_REPL.items():
        s = s.replace(a, b)
    return re.sub(r"[ \t]{2,}", " ", _RP_EMOJI.sub("", s)).strip()


def _rp_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _rp_field(par, instr, size=8):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(size)
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        elif kind == "text":
            r.text = txt
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def _rp_widths(df, total=17.4, minw=1.2):
    """Column widths (cm) proportional to the longest text in each column, summing to the page width."""
    w = []
    for c in df.columns:
        longest = max([len(w_) for w_ in _rp_clean(c).split()] + [0])
        cells = [min(len(_rp_clean(v)), 45) for v in df[c].head(60)]
        if max(cells + [0]) <= 3 and longest <= 5:  # rank / index columns
            w.append(3.5)
        else:
            w.append(max(7.5, longest * 1.1, (sum(cells) / max(len(cells), 1)) * 0.8, max(cells + [0]) * 0.45))
    tot = sum(w)
    ws = [max(minw, total * x / tot) for x in w]
    k = total / sum(ws)
    return [x * k for x in ws]


class _Rep:
    """Thin wrapper around python-docx with the handful of building blocks the report needs."""

    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
        nrm = self.d.styles["Normal"]
        nrm.font.name, nrm.font.size = "Calibri", Pt(10)
        nrm.paragraph_format.space_after = Pt(4)
        for name, size, color in (("Heading 1", 17, "1F3864"), ("Heading 2", 13, "2F5496"), ("Heading 3", 11, "2F5496")):
            s = self.d.styles[name]
            s.font.name, s.font.size, s.font.bold = "Calibri", Pt(size), True
            s.font.color.rgb = RGBColor.from_string(color)
            rf = s.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            s.paragraph_format.keep_with_next = True
        self.d.styles["Heading 1"].paragraph_format.page_break_before = True
        self.d.styles["Heading 1"].paragraph_format.space_after = Pt(8)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Oil Macro Environment Analyzer · educational only, not financial advice · page ")
        r.font.size = Pt(8)
        _rp_field(fp, "PAGE")

    # -- text
    def _runs(self, par, text, size=None, color=None, italic=False, bold=False):
        for i, part in enumerate(_rp_clean(text).split("**")):
            if not part:
                continue
            r = par.add_run(part)
            r.bold, r.italic = bold or i % 2 == 1, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def h1(self, t):
        self.d.add_heading(_rp_clean(t), 1)

    def h2(self, t):
        self.d.add_heading(_rp_clean(t), 2)

    def h3(self, t):
        self.d.add_heading(_rp_clean(t), 3)

    def p(self, t, **kw):
        par = self.d.add_paragraph()
        self._runs(par, t, **kw)
        return par

    def note(self, t):
        return self.p(t, size=8.5, color="595959", italic=True)

    def bullets(self, items):
        for it in items:
            self._runs(self.d.add_paragraph(style="List Bullet"), it)

    def box(self, t, kind="info"):
        fill = {"ok": "E2F0D9", "warn": "FFF2CC", "info": "DDEBF7", "bad": "F8D7DA"}[kind]
        tb = self.d.add_table(rows=1, cols=1)
        tb.style = "Table Grid"
        c = tb.rows[0].cells[0]
        c.width = Cm(17.4)
        _rp_shade(c, fill)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        self._runs(c.paragraphs[0], t, size=9.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- tables and charts
    def table(self, df, font=None, widths=None):
        if df is None or len(df) == 0:
            self.note("No data for this table.")
            return
        df = df.reset_index(drop=True)
        n = len(df.columns)
        font = font or (8 if n <= 6 else (7.5 if n <= 8 else 7))
        tb = self.d.add_table(rows=1, cols=n)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(df.columns):
            cell = tb.rows[0].cells[i]
            _rp_shade(cell, "1F3864")
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            self._runs(par, str(c), size=font, color="FFFFFF", bold=True)
        trPr = tb.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        for ri, row in enumerate(df.itertuples(index=False)):
            cells = tb.add_row().cells
            cs = OxmlElement("w:cantSplit")
            tb.rows[-1]._tr.get_or_add_trPr().append(cs)
            for i, v in enumerate(row):
                if ri % 2:
                    _rp_shade(cells[i], "F2F2F2")
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                self._runs(par, "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v), size=font)
        widths = widths or _rp_widths(df)
        tb.autofit = False
        tblPr = tb._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for i, w in enumerate(widths):
            tb.columns[i].width = Cm(w)
        for row in tb.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def kv(self, pairs, font=8.5):
        self.table(pd.DataFrame(pairs, columns=["Item", "Value"]), font=font, widths=[5.5, 11.9])

    def image(self, fig, title=None, w=17.2, h_px=520):
        if not _RP_IMG_OK[0] or fig is None:
            return False
        try:
            f2 = go.Figure(fig)
            f2.update_layout(template="plotly_white", paper_bgcolor="white", plot_bgcolor="white", title=title or (f2.layout.title.text if f2.layout.title else None))
            png = f2.to_image(format="png", width=1100, height=h_px, scale=2)
            self.d.add_picture(io.BytesIO(png), width=Cm(w))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
        except Exception:  # noqa: BLE001  (kaleido missing or failed: the report simply has no charts)
            _RP_IMG_OK[0] = False
            return False

    def save(self):
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()



def _rp_v(name, default=None):
    return globals().get(name, default)


def _rp_rank_rows():
    rows = []
    for c in chosen:
        xl, yl, xt, yt = F_ok[c].reindex(learn_idx), fwd.reindex(learn_idx), F_ok[c].reindex(unseen_idx), fwd.reindex(unseen_idx)
        dl = pd.concat([xl, yl], axis=1).dropna()
        t_l = nw_t(dl.iloc[:, 1].rank().values, dl.iloc[:, 0].rank().values, h) if len(dl) > 8 else np.nan
        rows.append(dict(c=c, ic_l=rank_ic(xl, yl), ic_t=rank_ic(xt, yt), t=t_l))
    R_ = pd.DataFrame(rows)
    R_["q"] = bh_q(R_["t"].map(lambda t: math.erfc(abs(t) / math.sqrt(2)) if pd.notna(t) else 1.0).values)
    return R_.sort_values("t", key=lambda s: -s.abs())


def _rp_oos_spreads():
    out = {}
    for nm, vv in (("Single-fit rules", verdict), ("Expanding rules", wfr["verdict"])):
        d = summarize(fwd, vv, grade_idx, h).set_index("Group")
        out[nm] = (d, (d.loc[FAV, "Avg"] - d.loc[UNF, "Avg"]) * 100 if d.loc[FAV, "Months"] >= 3 and d.loc[UNF, "Months"] >= 3 else np.nan)
    return out


def _rp_bt_calc():
    veh_ = "uso" if vehicle.startswith("USO") else "oil_fut"
    if veh_ not in m or not m[veh_].notna().any():
        return None
    px_ = m[veh_]
    rf_ = ((1 + m["irx"].fillna(0) / 100) ** (1 / 12) - 1) if "irx" in m else pd.Series(0.0, index=m.index)
    idx_ = grade_idx.intersection(px_.dropna().index)
    if len(idx_) < 12:
        return None
    nm_ = "USO" if veh_ == "uso" else "CL=F"
    r_s, w_s = run_backtest(px_, rf_, strategy_weights(strat, wfr["verdict"], px_ > px_.rolling(10).mean()), idx_, bps)
    r_b, _ = run_backtest(px_, rf_, pd.Series(1.0, index=px_.index), idx_, 0)
    cm = r_s.index.intersection(r_b.index)
    rows_ = [{"Portfolio": n_, "Months": (p_ := perf(r_.reindex(cm), rf_))["Months"], "CAGR": pc(p_["CAGR"]), "Volatility": pc(p_["Vol"], False),
              "Sharpe": "-" if pd.isna(p_["Sharpe"]) else f"{p_['Sharpe']:.2f}", "Max drawdown": pc(p_["MaxDD"])}
             for n_, r_ in ((f"Strategy: {strat}", r_s), (f"Buy & hold {nm_}", r_b))]
    fg = go.Figure()
    for n_, r_ in (("Strategy", r_s), (f"Buy & hold {nm_}", r_b)):
        fg.add_scatter(x=cm, y=(1 + r_.reindex(cm)).cumprod().values, name=n_)
    fg.update_layout(title=f"Growth of 1 on the {grade_name.lower()} months", legend=dict(orientation="h", y=-0.2))
    return pd.DataFrame(rows_), fg, w_s.reindex(cm).mean(), nm_, cm


def _rp_firing():
    pill, fired = {}, []
    for rules, sgn in ((good, 1), (bad, -1)):
        for (c, s_), w in rules.items():
            if S.loc[now, c] == s_:
                pill[META_ALL[c][4]] = pill.get(META_ALL[c][4], 0.0) + sgn * w
                fired.append({"Pillar": META_ALL[c][4], "Indicator": META_ALL[c][0], "Now": fmt_val(c, F_ok.loc[now, c]),
                              "State": state_label(c, s_), "Effect on oil": "Favorable" if sgn > 0 else "Unfavorable", "Weight": f"{w:.1f}"})
    return pill, fired


def _rp_summary(R):
    sp_ = _rp_oos_spreads()
    wv = wfr["verdict"].dropna()
    ex_ = wv.iloc[-1] if len(wv) else "n/a"
    R.box(f"**Single-fit rules: {v_now}** (score {sc_now:+.1f}). **Expanding rules (out-of-sample): {ex_}.** Historically, in environments like today's, WTI's {h}-month outcome "
          f"({target_name.lower()}) was {'better' if v_now == FAV else ('worse' if v_now == UNF else 'about average')} than usual. This describes the past; it is not a forecast.",
          "ok" if v_now == FAV else ("bad" if v_now == UNF else "info"))
    its = []
    for nm, (d, spd) in sp_.items():
        its.append(f"{nm} on the {grade_name.lower()} period: " + ("too few Favorable / Unfavorable months to compare." if pd.isna(spd) else f"Favorable minus Unfavorable = {spd:+.1f} pts over {h} months."))
    pill, fired = _rp_firing()
    if fired:
        top = sorted(fired, key=lambda x: -float(x["Weight"]))[:4]
        its.append("Strongest rules firing now: " + "; ".join(f"{x['Indicator']} ({x['State'].lower()}, {x['Effect on oil'].lower()})" for x in top) + ".")
    ok_exp = F_full[[c for c in EXP_KEYS if c in F_full]].dropna(how="all")
    if len(ok_exp):
        l_ = ok_exp.iloc[-1]
        if pd.notna(l_.get("ex_recess")):
            its.append(f"Yield-curve recession probability (12M): {l_['ex_recess']:.0f}%.")
        if pd.notna(l_.get("ex_ff_chg")):
            its.append(f"Curve-implied change in Fed funds over 12M: {l_['ex_ff_chg']:+.2f} pts.")
    if "ovx" in F_full and F_full["ovx"].notna().any():
        its.append(f"OVX (oil implied volatility): {F_full['ovx'].dropna().iloc[-1]:.1f}.")
    R.bullets(its)
    R.box("Oil is driven by supply shocks that macro indicators cannot see coming. Judge the signal by the out-of-sample and backtest sections, not by the in-sample environment table. "
          "Educational only, not financial advice.", "warn")


def _rp_settings(R):
    R.kv([("Data through", f"{m.index[-1]:%B %Y}"), ("Start date / usable months", f"{start} / {len(F_ok)}"), ("Prediction target", f"{target_name}, {h} months ahead"),
          ("Mode", "locked evaluation, final test revealed" if REVEAL else "research, final test hidden"),
          ("Learn / validation / final test", f"{learn_idx[0]:%b %Y}-{learn_idx[-1]:%b %Y} / {val_idx[0]:%b %Y}-{val_idx[-1]:%b %Y} / {test_idx[0]:%b %Y}-{test_idx[-1]:%b %Y}"
           if REVEAL else f"{learn_idx[0]:%b %Y}-{learn_idx[-1]:%b %Y} / {val_idx[0]:%b %Y}-{val_idx[-1]:%b %Y} / hidden"),
          ("States", "point-in-time percentile" if pit else "fixed terciles from the learn period"),
          ("Rule selection", f"{'FDR q <= ' if cfg[0] == 'q' else '|HAC t| >= '}{cfg[1]}, both halves required: {cfg[2]}, weighting: {cfg[3]}"),
          ("Backtest", f"{strat}; {vehicle.split(' (')[0]}; {bps} bps"), ("Indicators in use", str(len(chosen))),
          ("Expectation indicators in rules", "yes" if use_exp else "no")])
    miss_ = [k for k in failed if k != "eia_key"]
    R.p("Series that failed to download: " + (", ".join(miss_) if miss_ else "none") + (". No EIA key: OPEC, world balance and true futures curve are missing." if "eia_key" in failed else "."))
    R.p("Indicators in use: " + "; ".join(META_ALL[c][0] for c in chosen) + ".", size=9)


def _rp_dash(R):
    R.kv([("Single-fit verdict", f"{v_now} (score {sc_now:+.1f})"), ("Expanding-rules verdict", str(wfr['verdict'].dropna().iloc[-1]) if wfr['verdict'].notna().any() else "n/a"),
          ("WTI, latest month-end", f"${m['oil'].iloc[-1]:.1f}"), ("Rules in force", f"{len(good)} good / {len(bad)} bad")])
    f_ = go.Figure()
    f_.add_scatter(x=m["oil"].loc[start:].index, y=m["oil"].loc[start:].values, name="WTI", line=dict(color="#555", width=1.5))
    for v_, col_ in ((FAV, "#2e8b57"), (NEU, "#999999"), (UNF, "#c0392b")):
        ix = verdict.index[verdict == v_]
        f_.add_scatter(x=ix, y=m["oil"].reindex(ix).values, mode="markers", name=v_, marker=dict(color=col_, size=5))
    f_.update_layout(yaxis_type="log", legend=dict(orientation="h", y=-0.2))
    R.image(f_, "WTI with the environment verdict (single-fit rules)", h_px=440)
    pill, fired = _rp_firing()
    if pill:
        R.h2("Net score by pillar (rules firing today)")
        R.table(pd.DataFrame({"Pillar": list(pill), "Net score": [f"{v:+.1f}" for v in pill.values()]}).sort_values("Net score"))
    if fired:
        R.h2("Rules firing today")
        R.table(pd.DataFrame(fired))
    R.h2("All indicators now")
    R.table(pd.DataFrame([{"Pillar": META_ALL[c][4], "Indicator": META_ALL[c][0], "Now": fmt_val(c, F_ok.loc[now, c]),
                           "Percentile": f"{P_full[c].dropna().iloc[-1]:.0%}" if P_full[c].notna().any() else "-",
                           "Reads as": state_label(c, S.loc[now, c]) if pd.notna(S.loc[now, c]) else "-"} for c in chosen]).sort_values(["Pillar", "Indicator"]))


def _rp_guide(R):
    R.p("Question: which macro environments were historically good or bad for WTI over the next months, and does that survive on unseen data?")
    R.bullets([f"**{p}**: " + ", ".join(META_ALL[c][0] for c in chosen if META_ALL[c][4] == p) for p in PILLARS if any(META_ALL[c][4] == p for c in chosen)])
    R.p("Method: every indicator is dated to when it was public and labelled Low / Mid / High. Rules are fitted on the learn months (HAC t-stats for overlapping windows, "
        f"FDR q-values, effect required in both halves), chosen on validation and graded once on the hidden final test. The expanding rules are refitted every {RULE_STEP} months using only outcomes already known.")
    R.note("Oil-specific cautions: supply shocks are unpredictable; inventories use a seasonal norm; spot returns are not futures returns (roll yield); 2008 and 2020 are single episodes that can dominate any statistic.")


def _rp_env(R):
    fdd_ = fdd if target in ("ret", "vol") else None
    for nm, idx, vv in (("Learn period (in-sample)", learn_idx, verdict), (f"{grade_name}, single-fit rules (unseen)", grade_idx, verdict),
                        (f"{grade_name}, expanding rules (unseen)", grade_idx, wfr["verdict"])):
        R.h2(nm)
        R.table(fmt_summary(summarize(fwd, vv, idx, h, fdd_), dd=fdd_ is not None))
    R.note("Learn months are in-sample by construction. The grading months are the real test.")


def _rp_rank(R):
    out = []
    for r in _rp_rank_rows().itertuples():
        same = bool(pd.notna(r.ic_l) and pd.notna(r.ic_t) and np.sign(r.ic_l) == np.sign(r.ic_t))
        status = "-" if (pd.isna(r.ic_l) or pd.isna(r.ic_t)) else ("Flips on unseen data" if not same else ("Consistent and significant" if abs(r.t) >= 2 else "Consistent but weak"))
        out.append({"Indicator": META_ALL[r.c][0], "IC learn": "-" if pd.isna(r.ic_l) else f"{r.ic_l:+.2f}", f"IC {unseen_lbl}": "-" if pd.isna(r.ic_t) else f"{r.ic_t:+.2f}",
                    "HAC t": "-" if pd.isna(r.t) else f"{r.t:+.1f}", "FDR q": f"{r.q:.2f}", "Status": status})
    R.table(pd.DataFrame(out).head(25))
    R.note(f"Top 25 of {len(out)} indicators by |t|. About 5% of indicators look significant by luck alone; prefer those that keep their sign on unseen data and have an economic story.")


def _rp_oos(R):
    for nm, (d, spd) in _rp_oos_spreads().items():
        R.h2(nm)
        R.table(fmt_summary(d.reset_index()))
        if pd.notna(spd):
            R.box(f"Favorable minus Unfavorable = {spd:+.1f} pts over {h} months on the {grade_name.lower()} period.", "ok" if spd > 0 else "warn")
    R.note(f"About {len(grade_idx)} graded months, but {h}-month windows overlap, so independent observations are roughly {max(1, len(grade_idx) // h)}.")


def _rp_backtest(R, charts):
    bt = _rp_bt_calc()
    if bt is None:
        R.p("The backtest could not run (vehicle prices unavailable or fewer than 12 graded months with prices).")
        return
    tb, fg, expo, nm_, cm = bt
    R.p(f"Expanding-rules signal, vehicle {nm_}, {bps} bps per 100% traded, graded months {cm[0]:%b %Y} to {cm[-1]:%b %Y}. Idle cash earns T-bills. Average exposure {expo:.0%}.")
    R.table(tb)
    if charts:
        R.image(fg, h_px=420)
    if nm_ == "CL=F":
        R.box("CL=F is a stitched front-month series that ignores roll yield, so these results are not what a futures holder earned. USO is more realistic.", "warn")
    R.note("A short, single-regime sample. Compare with the trend-only strategy to see whether the macro rules add anything.")


def _rp_expectations(R):
    xc = [c for c in EXP_KEYS if c in F_full.columns]
    if not xc:
        R.p("No expectation indicator could be built (Treasury / FRED series missing).")
        return
    rows = []
    for c in xc:
        s_ = F_full[c].dropna()
        if s_.empty:
            continue
        prev = F_full[c].shift(3).loc[s_.index[-1]]
        rows.append({"Indicator": META_ALL[c][0], "Method": EXP_INFO[c]["method"], "Latest": _fx(c, s_.iloc[-1]),
                     "vs 3 months ago": "-" if pd.isna(prev) else _fx(c, s_.iloc[-1] - prev, True), "Textbook sign for oil": _STXT[OIL_EXP_SIGN[c]]})
    R.table(pd.DataFrame(rows))
    R.h2("Do the expectations line up with oil's outcome?")
    out = []
    for c in xc:
        x_ = F_full[c].reindex(fwd.index)
        ic_l, ic_t = rank_ic(x_.loc[learn_idx], fwd.loc[learn_idx]), rank_ic(x_.loc[unseen_idx], fwd.loc[unseen_idx])
        out.append({"Indicator": META_ALL[c][0], "IC learn": "-" if pd.isna(ic_l) else f"{ic_l:+.2f}", f"IC {unseen_lbl}": "-" if pd.isna(ic_t) else f"{ic_t:+.2f}",
                    "Same sign": "n/a" if pd.isna(ic_l) or pd.isna(ic_t) else ("Yes" if np.sign(ic_l) == np.sign(ic_t) else "No")})
    R.table(pd.DataFrame(out))
    gr = _grade(m, F_full[xc], ev_r)
    if not gr.empty:
        R.h2("Are the expectations any good? (12 months later)")
        R.table(pd.DataFrame({"Expectation": gr["k"].map(lambda k: META_ALL[k][0]), "Months": gr["n"], "Skill vs 'nothing changes'": gr["skill"].map(pc),
                              "IC expected vs realised": gr["ic"].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}")}))
        R.note("Skill below zero means 'nothing changes' was a better guess than the expectation.")


def _rp_options(R):
    R.p("Historical part (in the rules): OVX, the options-implied volatility of oil, and its gap to realized volatility. Live part: a snapshot of USO options (not backtestable, Yahoo has no option history).")
    pairs = []
    for c, nm in (("ovx", "OVX"), ("ovx_vrp", "Implied minus realized vol (pts)"), ("gpr", "Geopolitical risk index")):
        if c in F_full and F_full[c].notna().any():
            pairs.append((nm, f"{F_full[c].dropna().iloc[-1]:.1f}"))
    if pairs:
        R.kv(pairs)
    sn = _rp_v("snap")
    if sn:
        R.h2("USO options snapshot")
        R.kv([("Spot", f"{sn['spot']:.2f}"), (f"ATM implied vol, {sn['near_dte']}d expiry", f"{sn['atm_near']:.1f}%"),
              ("ATM implied vol, far expiry", "-" if pd.isna(sn["atm_far"]) else f"{sn['atm_far']:.1f}%"), ("Term (far minus near)", "-" if pd.isna(sn["term"]) else f"{sn['term']:+.1f} pts"),
              ("Put skew minus call skew", "-" if pd.isna(sn["skew"]) else f"{sn['skew']:+.1f} pts"), ("Put/call ratio (volume / open interest)", f"{sn['pc_vol']:.2f} / {sn['pc_oi']:.2f}")])
        R.bullets(_opt_read(sn))
    else:
        R.p("The live options snapshot was not available when the report was built.")
    R.note("USO options carry ETF and roll effects and are not the WTI futures options market. Treat as context, not a signal.")


def _rp_data(R):
    rows = []
    for k, (desc, src) in SOURCES.items():
        ok = k in m and m[k].notna().any()
        rows.append({"Series": desc, "Source": src, "Status": "OK" if ok else "Missing", "First": f"{m[k].dropna().index[0]:%b %Y}" if ok else "-"})
    R.table(pd.DataFrame(rows))
    R.note("Weekly EIA data is used from the Wednesday release; EIA STEO, OECD, GPR and macro series are current-vintage (mild look-ahead from revisions).")


def build_word_report(charts=True):
    """Returns (docx bytes, list of section names that could not be built)."""
    _RP_IMG_OK[0] = bool(charts)
    R, fails = _Rep(), []
    t = R.d.add_paragraph()
    R._runs(t, "Oil Macro Environment Report", size=26, color="1F3864", bold=True)
    R.p(f"Results of all tabs · data through {m.index[-1]:%B %Y} · prepared {pd.Timestamp.today():%d %B %Y}", size=11, color="595959")
    R.p(f"Predicting: {target_name}, {h} months ahead · mode: {'locked evaluation, final test revealed' if REVEAL else 'research (final test hidden)'}", size=10, color="595959")
    plan = [("Executive summary", _rp_summary), ("Settings and data", _rp_settings), ("Dashboard", _rp_dash), ("Guide: how this works", _rp_guide),
            ("Environments", _rp_env), ("Indicator ranking", _rp_rank), ("Out-of-sample", _rp_oos), ("Backtest", lambda R_: _rp_backtest(R_, charts)),
            ("Expectations", _rp_expectations), ("Options & geopolitics", _rp_options), ("Data & coverage", _rp_data)]
    for i, (name, fn) in enumerate(plan, 1):
        R.h1(f"{i}. {name}")
        if i == 1:
            R.d.paragraphs[-1].paragraph_format.page_break_before = False
        try:
            fn(R)
        except Exception as e:  # noqa: BLE001  (one broken section must not lose the whole report)
            fails.append(f"{name} ({type(e).__name__}: {e})")
            R.box(f"This section could not be built ({type(e).__name__}: {e}). The other sections are unaffected.", "bad")
    return R.save(), fails


with T["report"]:
    st.subheader("📄 Word report: the results of all tabs in one document")
    st.caption("Builds a Word (.docx) file with an executive summary, the settings used and the results of every tab: tables, verdicts and the main charts. It reports exactly what the app shows "
               "for the current settings, including the research / locked mode: in research mode the final test stays hidden. Press the button after you have finished changing settings; "
               "the report is not rebuilt on every click.")
    if not _HAS_DOCX:
        st.error("The python-docx package is not installed. Add `python-docx` to requirements.txt (and optionally `kaleido` for charts), then restart the app.")
    else:
        want_charts = st.checkbox("Include charts (needs the optional `kaleido` package; without it the report is built without charts)", True, key="oil_rep_charts")
        sig = repr((str(m.index[-1]), start, target, h, REVEAL, sorted(chosen), pit, cfg, use_exp, strat, vehicle, bps, want_charts))
        if st.button("📝 Build the Word report", key="oil_rep_build"):
            with st.spinner("Collecting results from every tab and writing the Word file..."):
                _data, _fails = build_word_report(want_charts)
            st.session_state["oil_rep_saved"] = (sig, _data, _fails, pd.Timestamp.now())
        saved = st.session_state.get("oil_rep_saved")
        if saved:
            sig0, data_, fails_, ts_ = saved
            if sig0 != sig:
                st.warning("Settings or data have changed since this report was built. Rebuild it to match what the tabs show now.")
            st.success(f"Report ready ({len(data_) / 1024:,.0f} KB, built {ts_:%H:%M:%S}).")
            st.download_button("⬇️ Download the report (.docx)", data_, file_name=f"oil_report_{now:%Y_%m}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="oil_rep_dl")
            if fails_:
                st.warning("Some sections could not be built and contain an error note instead: " + "; ".join(fails_))
            if want_charts and not _RP_IMG_OK[0]:
                st.info("Charts were left out because the `kaleido` package is not installed. Add `kaleido` to requirements.txt to include them.")
