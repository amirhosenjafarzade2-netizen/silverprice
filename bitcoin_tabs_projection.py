"""
Bitcoin tabs, part 4: Projection (ML route fan, now with optional expectation features), Machine learning, Look-ahead scan, Bitcoin & gold now.
Executed by bitcoin_main.py (all config, helpers and results are available as globals).
"""
import math

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from scipy.special import ndtri
from sklearn.ensemble import GradientBoostingRegressor
from sklearn.linear_model import Ridge
from sklearn.metrics import roc_auc_score
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler


# results of the projection tab, read by the Word report tab at the end of this file
PROJ_REPORT = {}


# ------------------------------------------------------------------ projection (ML route fan)
def proj_rank_ic(a, b):
    d = pd.concat([pd.Series(np.asarray(a, float)), pd.Series(np.asarray(b, float))], axis=1).dropna()
    if len(d) < 8 or d.iloc[:, 0].nunique() < 2:
        return np.nan
    return d.iloc[:, 0].rank().corr(d.iloc[:, 1].rank())


def pinball(y, Q, qs=PROJ_Q):
    e = np.asarray(y)[:, None] - Q
    qa = np.asarray(qs)
    return float(np.maximum(qa * e, (qa - 1) * e).mean())


def proj_price_features(m, idx):
    s = m["btc"]
    f = pd.DataFrame(index=m.index)
    f["px_mom1"], f["px_mom3"] = s.pct_change(1), s.pct_change(3)
    f["px_mom6"], f["px_mom12"] = s.pct_change(6), s.pct_change(12)
    f["px_vol12"] = s.pct_change().rolling(12).std() * np.sqrt(12)
    f["px_trend"] = s / s.rolling(10).mean() - 1
    f["px_dd12"] = s / s.rolling(12).max() - 1
    return f.reindex(idx)


def proj_inputs(m, F_ok, feats, length, E_exp=None):
    """Features known at each month-end and forward log returns for every grid horizon up to `length`.
    E_exp = frame of the expectation indicators (point-in-time). PROJ_FEATS[3] adds those not already in F_ok; PROJ_FEATS[4] uses all of them with bitcoin momentum."""
    pf = proj_price_features(m, F_ok.index)
    ex_all = E_exp.reindex(F_ok.index) if E_exp is not None and len(E_exp.columns) else pd.DataFrame(index=F_ok.index)
    ex_new = ex_all[[c for c in ex_all.columns if c not in F_ok.columns]]
    if feats == PROJ_FEATS[1]:
        X = F_ok
    elif feats == PROJ_FEATS[2]:
        X = pf
    elif feats == PROJ_FEATS[3]:
        X = pd.concat([F_ok, pf, ex_new], axis=1)
    elif feats == PROJ_FEATS[4]:
        X = pd.concat([ex_all, pf], axis=1)
    else:
        X = pd.concat([F_ok, pf], axis=1)
    X = X.replace([np.inf, -np.inf], np.nan).dropna()
    ks = [k for k in PROJ_GRID if k <= length]
    s = m["btc"]
    Y = pd.DataFrame({k: np.log(s.shift(-k) / s) for k in ks}).reindex(X.index)
    return X, Y, ks


def _fit_quantiles(model_name, Xtr, ytr, qs, Xte):
    """Predicted quantiles (rows = Xte, columns = qs)."""
    if model_name == PJ_RIDGE:
        mdl = make_pipeline(StandardScaler(), Ridge(alpha=PROJ_ALPHA)).fit(Xtr, ytr)
        rq = np.quantile(ytr - mdl.predict(Xtr), qs)
        return mdl.predict(Xte)[:, None] + rq[None, :]
    cols = [GradientBoostingRegressor(loss="quantile", alpha=q, n_estimators=40, max_depth=2, learning_rate=0.05, subsample=0.8,
                                      min_samples_leaf=10, random_state=0).fit(Xtr, ytr).predict(Xte) for q in qs]
    return np.sort(np.column_stack(cols), axis=1)


@st.cache_data(show_spinner=False)
def proj_walk_forward(X, Y, qs, cut, end, step, model_name):
    """Every `step` months refit using only rows whose k-month outcome was already known, predict the next `step` months.
    Returns {k: (dates, model quantiles, history-only quantiles, realised log return)}. No Streamlit calls inside."""
    Xv, idx, out = X.values.astype(float), X.index, {}
    for k in Y.columns:
        yv = Y[k].values.astype(float)
        dts, M, B, yy = [], [], [], []
        for s0 in range(cut, end, step):
            te = np.arange(s0, min(s0 + step, end))
            tr = np.arange(0, max(s0 - int(k) + 1, 0))
            tr = tr[~np.isnan(yv[tr])]
            if len(tr) < PROJ_MINTR or len(te) == 0:
                continue
            M.append(_fit_quantiles(model_name, Xv[tr], yv[tr], qs, Xv[te]))
            B.append(np.tile(np.quantile(yv[tr], qs), (len(te), 1)))
            dts.extend(idx[te])
            yy.extend(yv[te])
        if M:
            out[int(k)] = (pd.DatetimeIndex(dts), np.vstack(M), np.vstack(B), np.array(yy, float))
    return out


@st.cache_data(show_spinner=False)
def proj_fit_today(X, Y, qs, model_name):
    """Fit on every row whose outcome is known and predict from the latest month. Returns (horizons, model quantiles, history quantiles)."""
    Xv = X.values.astype(float)
    ks, Qm, Qb = [], [], []
    for k in Y.columns:
        yv = Y[k].values.astype(float)
        tr = np.where(~np.isnan(yv))[0]
        if len(tr) < PROJ_MINTR:
            continue
        Qm.append(_fit_quantiles(model_name, Xv[tr], yv[tr], qs, Xv[-1:])[0])
        Qb.append(np.quantile(yv[tr], qs))
        ks.append(int(k))
    return ks, np.array(Qm), np.array(Qb)


def _proj_members(model_name):
    return [PJ_RIDGE, PJ_GBR] if model_name == PJ_AVG else [model_name]


def _combine_wf(parts):
    if len(parts) == 1:
        return parts[0]
    out = {}
    for k in parts[0]:
        if all(k in p for p in parts):
            idx, _, B, y = parts[0][k]
            out[k] = (idx, np.mean([p[k][1] for p in parts], axis=0), B, y)
    return out


def to_curve(ks, Q, L):
    """Interpolate per-horizon quantiles to every month 0..L (0 at month 0)."""
    mo = np.arange(0, L + 1)
    return np.column_stack([np.interp(mo, [0] + list(ks), [0.0] + list(Q[:, j])) for j in range(Q.shape[1])])


def proj_paths(Qc, qs, S0, n=PROJ_PATHS, seed=7):
    """Routes whose month-j value follows the predicted quantiles of month j exactly. A standardised random walk (W_j / sqrt(j) ~ N(0,1))
    supplies the month-to-month dependence; each month's z-score is mapped through that month's quantile curve (linear tails)."""
    L = Qc.shape[0] - 1
    zq = ndtri(np.asarray(qs))
    rng = np.random.default_rng(seed)
    e = rng.standard_normal((n // 2, L))
    e = np.vstack([e, -e])  # antithetic pairs
    Z = np.cumsum(e, axis=1) / np.sqrt(np.arange(1, L + 1))
    lr = np.zeros((len(e), L + 1))
    for j in range(1, L + 1):
        v = np.maximum.accumulate(Qc[j])
        slo = max((v[1] - v[0]) / (zq[1] - zq[0]), 1e-6)
        shi = max((v[-1] - v[-2]) / (zq[-1] - zq[-2]), 1e-6)
        z = Z[:, j - 1]
        y = np.interp(z, zq, v)
        y = np.where(z < zq[0], v[0] + slo * (z - zq[0]), y)
        y = np.where(z > zq[-1], v[-1] + shi * (z - zq[-1]), y)
        lr[:, j] = y
    return S0 * np.exp(lr)


def proj_score(wfd, lam, eval_idx, ks):
    """Walk-forward grade on unseen months: model fan (shrunk by lam toward history) versus the history-only fan."""
    rows, mid = [], len(PROJ_Q) // 2
    for k in ks:
        if k not in wfd:
            continue
        idx, M, B, y = wfd[k]
        ok = np.asarray(idx.isin(eval_idx)) & ~np.isnan(y)
        if ok.sum() < 8:
            continue
        Qb = B[ok]
        Qm = np.sort(Qb + lam * (M[ok] - Qb), axis=1)
        yy = y[ok]
        pm, pb = pinball(yy, Qm), pinball(yy, Qb)
        call, real = np.sign(Qm[:, mid] - Qb[:, mid]), np.sign(yy - Qb[:, mid])
        nz = call != 0
        rows.append(dict(k=k, n=int(ok.sum()), ic=proj_rank_ic(Qm[:, mid], yy),
                         hit=float((call[nz] == real[nz]).mean()) if nz.any() else np.nan,
                         skill=1 - pm / pb if pb > 0 else np.nan,
                         c80=float(((yy >= Qm[:, 0]) & (yy <= Qm[:, -1])).mean()),
                         c50=float(((yy >= Qm[:, 1]) & (yy <= Qm[:, -2])).mean())))
    return pd.DataFrame(rows)


def compare_expectations(m, F_ok, E_exp, length, model_name, val_idx, reveal, ev_, lam):
    """Same model, same unseen months, with and without the expectation indicators. Diagnostic only, not a menu to pick a winner from."""
    ex_cols = [c for c in E_exp.columns if c in F_ok.columns]
    F_base = F_ok.drop(columns=ex_cols)  # if the sidebar switch already put them in F_ok, take them out of the 'without' arm
    arms = {"Macro + bitcoin momentum (without expectations)": (F_base, PROJ_FEATS[0]),
            "Macro + bitcoin momentum + expectations": (F_base, PROJ_FEATS[3])}
    built = {k: proj_inputs(m, fb, ft, length, E_exp) for k, (fb, ft) in arms.items()}
    common = pd.DatetimeIndex(ev_)
    for X, _, _ in built.values():
        common = common.intersection(X.index)
    members = _proj_members(model_name)
    rows = []
    for name, (X, Y, ks) in built.items():
        if len(X) < 120 or not ks:
            continue
        cut = int(X.index.searchsorted(val_idx[0]))
        end = len(X) if reveal else int(X.index.searchsorted(val_idx[-1], side="right"))
        wf_ = _combine_wf([proj_walk_forward(X, Y, PROJ_Q, cut, end, PROJ_STEP, nm) for nm in members])
        sc_ = proj_score(wf_, lam, common, ks)
        if sc_.empty:
            continue
        rows.append({"Feature set": name, "Features": X.shape[1], "Unseen months (smallest horizon)": int(sc_["n"].min()),
                     "Average skill vs history": pc(sc_["skill"].mean()), "_skill": float(sc_["skill"].mean()), "Average rank IC": f"{sc_['ic'].mean():+.2f}",
                     "Horizons better than history": f"{int(((sc_['skill'] > 0) & (sc_['ic'] > 0)).sum())} of {len(sc_)}",
                     "Average 80% coverage": pc(sc_["c80"].mean(), False)})
    return pd.DataFrame(rows)


def render_projection_tab(m, F_ok, val_idx, unseen_idx, unseen_lbl, reveal, model_name, feats, length, trust, label, E_exp=None):
    """label(col) -> display name of a feature column. E_exp = expectation indicators (may be None / empty)."""
    PROJ_REPORT.clear()
    lam = trust / 100.0
    st.subheader(f"🔮 Bitcoin price projection: the next {length} months")
    st.caption("A machine-learning FAN of possible routes, not a single forecast line. Each month's spread is the model's predicted distribution of bitcoin's return "
               "(10 / 25 / 50 / 75 / 90% levels), shrunk toward the unconditional history by the trust setting; sample routes move like a random walk but "
               "match that distribution at every month. The grading table below checks, on months the model never saw, whether it beats a history-only fan. "
               "Today's fit uses every outcome known so far (like the other 'now' reads); grading uses only unseen months.")

    if feats in (PROJ_FEATS[3], PROJ_FEATS[4]) and (E_exp is None or not len(E_exp.columns)):
        st.warning("No expectation indicators are available (their data did not download), so this feature set falls back to "
                   + ("macro + bitcoin momentum." if feats == PROJ_FEATS[3] else "bitcoin momentum only."))
    X, Y, ks_all = proj_inputs(m, F_ok, feats, length, E_exp)
    if len(X) < 120 or not ks_all:
        st.info("Not enough history for the projection with these settings. Move the start date earlier or use more indicators.")
        return
    members = _proj_members(model_name)
    cut = int(X.index.searchsorted(val_idx[0]))
    end = len(X) if reveal else int(X.index.searchsorted(val_idx[-1], side="right"))
    with st.spinner("Fitting projection models (cached after the first run)..."):
        wf = _combine_wf([proj_walk_forward(X, Y, PROJ_Q, cut, end, PROJ_STEP, nm) for nm in members])
        fits = [proj_fit_today(X, Y, PROJ_Q, nm) for nm in members]
    ks, Qm, Qb = fits[0][0], np.mean([f[1] for f in fits], axis=0), fits[0][2]
    if not ks:
        st.info("Not enough training rows to fit any horizon.")
        return
    L = min(length, max(ks))
    Qs = np.sort(Qb + lam * (Qm - Qb), axis=1)
    Qc, Qbc = to_curve(ks, Qs, L), to_curve(ks, Qb, L)
    S0 = float(m["btc"].loc[X.index[-1]])
    P = proj_paths(Qc, PROJ_Q, S0)
    fut = pd.DatetimeIndex([X.index[-1] + pd.offsets.MonthEnd(j) for j in range(L + 1)])
    mid = len(PROJ_Q) // 2

    # --- headline numbers
    run_max = np.maximum.accumulate(P, axis=1)
    mdd = (P / run_max - 1).min(axis=1)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric(f"Median price in {L}m", f"${S0 * np.exp(Qc[L, mid]):,.1f}", f"{np.exp(Qc[L, mid]) - 1:+.1%} vs ${S0:,.1f} now")
    c2.metric(f"80% range in {L}m", f"${S0 * np.exp(Qc[L, 0]):,.0f} – ${S0 * np.exp(Qc[L, -1]):,.0f}")
    c3.metric(f"Chance higher in {L}m", f"{(P[:, L] > S0).mean():.0%}")
    c4.metric(f"Chance of a {PROJ_DD:.0%}+ drawdown on the way", f"{(mdd <= -PROJ_DD).mean():.0%}",
              f"typical worst fall {np.median(mdd):.0%}", delta_color="off")

    PROJ_REPORT.update(L=L, S0=S0, trust=trust, model=model_name, feats=feats, last=X.index[-1],
                       head=dict(median=S0 * np.exp(Qc[L, mid]), chg=np.exp(Qc[L, mid]) - 1, lo=S0 * np.exp(Qc[L, 0]), hi=S0 * np.exp(Qc[L, -1]),
                                 up=float((P[:, L] > S0).mean()), dd=float((mdd <= -PROJ_DD).mean()), dd_med=float(np.median(mdd))))

    # --- fan chart
    hist = m["btc"].loc[: X.index[-1]].iloc[-60:]
    f = go.Figure()
    f.add_scatter(x=hist.index, y=hist.values, name="Bitcoin (history)", line=dict(color="#555"))

    def band(lo, hi, color, name):
        f.add_scatter(x=list(fut) + list(fut[::-1]), y=list(S0 * np.exp(Qc[:, hi])) + list(S0 * np.exp(Qc[::-1, lo])),
                      fill="toself", fillcolor=color, line=dict(width=0), name=name, hoverinfo="skip")

    band(0, len(PROJ_Q) - 1, "rgba(74,111,165,0.15)", "10–90% range")
    band(1, len(PROJ_Q) - 2, "rgba(74,111,165,0.30)", "25–75% range")
    for i in range(25):
        f.add_scatter(x=fut, y=P[i], line=dict(width=1, color="rgba(70,110,170,0.22)"), showlegend=False, hoverinfo="skip")
    f.add_scatter(x=fut, y=S0 * np.exp(Qbc[:, mid]), name="History-only median", line=dict(color="#999", dash="dash", width=2))
    f.add_scatter(x=fut, y=S0 * np.exp(Qc[:, mid]), name="ML median route", line=dict(color="#1f5fbf", width=3))
    vmark(f, fut[0], "today")
    f.update_layout(title=f"Bitcoin route fan: {model_name}, trust {trust}%", yaxis_title="USD", yaxis_type="log", height=500)
    show_plot(st, f)
    PROJ_REPORT["fig"] = f
    st.caption("Shaded bands are the model's predicted distribution at each month; the thin lines are 25 of 1,000 sample routes; "
               "the dashed grey line is what plain history alone would say. The gap between the blue and grey medians is everything the model adds.")

    # --- table by horizon
    rows = []
    for k in sorted({k for k in (1, 3, 6, 12, 18, 24) if k <= L} | {L}):
        rows.append({"Horizon": f"{k} months", "Median price": f"${S0 * np.exp(Qc[k, mid]):,.1f}",
                     "Median change": pc(np.exp(Qc[k, mid]) - 1), "History-only median": pc(np.exp(Qbc[k, mid]) - 1),
                     "25–75% range": f"{np.exp(Qc[k, 1]) - 1:+.0%} to {np.exp(Qc[k, -2]) - 1:+.0%}",
                     "10–90% range": f"{np.exp(Qc[k, 0]) - 1:+.0%} to {np.exp(Qc[k, -1]) - 1:+.0%}",
                     "Chance higher": f"{(P[:, k] > S0).mean():.0%}"})
    show_df(st, pd.DataFrame(rows))
    PROJ_REPORT["table"] = pd.DataFrame(rows)

    # --- grading
    st.subheader("Does the model beat plain history? (walk-forward, unseen months only)")
    ev_ = pd.DatetimeIndex(unseen_idx)
    sc = proj_score(wf, lam, ev_, ks)
    st.caption(f"Graded on {unseen_lbl} months. Each refit uses only rows whose outcome was already known. "
               "**Skill** = reduction in pinball loss (the quantile-forecast error) versus the history-only fan: positive means the model's fan was better. "
               "**Coverage** should be near its nominal level (80% / 50%). **Rank IC** = rank correlation of the model's median with the realised return. "
               "**Direction hit** = how often the model's tilt above / below the history median matched the outcome (50% = coin flip). "
               "Windows overlap, so the number of independent observations is roughly months ÷ horizon.")
    if sc.empty:
        st.info("Not enough unseen months to grade the projection.")
    else:
        _sc_tab = pd.DataFrame({
            "Horizon": [f"{k} months" for k in sc["k"]], "Unseen months": sc["n"],
            "Skill vs history": sc["skill"].map(lambda v: pc(v)), "Rank IC": sc["ic"].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}"),
            "Direction hit": sc["hit"].map(lambda v: pc(v, False)),
            "80% range coverage": sc["c80"].map(lambda v: pc(v, False)), "50% range coverage": sc["c50"].map(lambda v: pc(v, False))})
        show_df(st, _sc_tab)
        good_ = (sc["skill"] > 0) & (sc["ic"] > 0)
        share, avg = float(good_.mean()), float(sc["skill"].mean())
        cov = float(sc["c80"].mean())
        if share >= 0.7 and avg > 0.02:
            st.success(f"The model beat history-only on {good_.sum()} of {len(sc)} horizons (average skill {avg:+.1%}). Some real information, but the test is short "
                       "and windows overlap: read the median tilt as a mild lean, not a target.")
        else:
            st.warning(f"The model did NOT reliably beat history-only ({good_.sum()} of {len(sc)} horizons better, average skill {avg:+.1%}). "
                       "Treat the fan as a risk range (how wide routes can be) and ignore any tilt in the median. Lowering the trust setting moves the fan toward history.")
        PROJ_REPORT.update(score=_sc_tab, unseen_lbl=unseen_lbl, n_good=int(good_.sum()), n_hor=len(sc), avg_skill=avg, share=share, cov=cov)
        if cov < 0.70 or cov > 0.90:
            st.warning(f"The 80% range contained the outcome {cov:.0%} of the time on unseen months, so the fan is {'too narrow' if cov < 0.70 else 'too wide'}.")

        tr_rows = []
        for l_ in (0.0, 0.25, 0.5, 0.75, 1.0):
            s_ = proj_score(wf, l_, ev_, ks)
            if not s_.empty:
                tr_rows.append({"Trust in model": f"{l_:.0%}", "Average skill vs history": pc(s_["skill"].mean()),
                                "Average rank IC": "-" if s_["ic"].isna().all() else f"{s_['ic'].mean():+.2f}",
                                "Average 80% coverage": pc(s_["c80"].mean(), False)})
        st.markdown("**Skill by trust level** (a diagnostic, not a menu: choosing the best row would select on the grading months)")
        show_df(st, pd.DataFrame(tr_rows))
        PROJ_REPORT["trust_tab"] = pd.DataFrame(tr_rows)

    # --- do the expectations help?
    if E_exp is not None and len(E_exp.columns):
        with st.expander("🧪 Do the expectation indicators help the projection? (second fit on the same unseen months)"):
            st.caption("Refits the same model twice, once on macro + bitcoin momentum and once with the expectation indicators added, and grades both on the same unseen months "
                       "(runs the walk-forward again, so it takes a moment the first time). A diagnostic, not a menu: the expectation indicators start later than most macro series, "
                       "so both arms use only the months where all inputs exist, and gains of a point or two of skill are within noise.")
            if st.checkbox("Run the with / without comparison", False, key="pj_cmp"):
                with st.spinner("Fitting both versions (cached after the first run)..."):
                    cmp_ = compare_expectations(m, F_ok, E_exp, length, model_name, val_idx, reveal, ev_, lam)
                if cmp_.empty:
                    st.info("Not enough overlapping history to compare the two versions.")
                else:
                    show_df(st, cmp_.drop(columns="_skill"))
                    PROJ_REPORT["compare"] = cmp_.drop(columns="_skill")
                    if len(cmp_) == 2:
                        d_sk = float(cmp_["_skill"].iloc[1] - cmp_["_skill"].iloc[0])
                        (st.success if d_sk > 0.01 else st.info)(
                            f"With expectations the average skill moved by {d_sk * 100:+.1f} percentage points versus the version without. "
                            + ("A small gain; check that it also shows in the rank IC and holds when you change the horizon length or the model." if d_sk > 0.01
                               else "No clear gain: the expectation indicators did not add reliable information to this projection on unseen months."))

    # --- drivers
    kd = min(ks, key=lambda k: abs(k - 6))
    ok_ = Y[kd].notna().values
    mdl = make_pipeline(StandardScaler(), Ridge(alpha=PROJ_ALPHA)).fit(X.values[ok_], Y[kd].values[ok_])
    scl, rg = mdl[0], mdl[-1]
    contrib = pd.Series(rg.coef_ * (X.iloc[-1].values - scl.mean_) / scl.scale_, index=X.columns) * 100
    contrib = contrib.reindex(contrib.abs().sort_values(ascending=False).index).head(12)
    st.subheader(f"What is pushing the {kd}-month projection today?")
    st.bar_chart(contrib.rename(index={c: label(c) for c in contrib.index}).sort_values())
    PROJ_REPORT.update(drivers=contrib.rename(index={c: label(c) for c in contrib.index}), kd=kd)
    st.caption(f"Contribution of each feature to the {kd}-month expected log return, in percentage points relative to an average month (ridge view, shown for every model). "
               "Correlated indicators share credit unpredictably, so read this as a story about the inputs, not as causes.")
    st.warning("**Not a price target and not financial advice.** The fan describes how wide bitcoin's routes have been after similar conditions, "
               "adjusted by a model that may have no skill. Shocks, news and regime changes are not in the data.")


# ------------------------------------------------------------------ projection tab
with T["proj"]:
    render_projection_tab(m, F_ok, val_idx, unseen_idx, unseen_lbl, REVEAL, proj_model, proj_feats, proj_len, proj_trust,
                          lambda c: META_ALL[c][0] if c in META_ALL else PROJ_NAMES.get(c, c),
                          F_full[[c for c in EXP_KEYS if c in F_full.columns]])

# ------------------------------------------------------------------ machine learning
if run_ml:
    with T["ml"]:
        st.markdown(f"""
Each model estimates the **probability that {tgt_txt}** in {h} months, retrained every {step} months on data whose outcome was already known at that time.
Scores below use only months whose outcome is known now. **AUC** 0.50 = coin flip, 0.55-0.60 = modest skill, above 0.65 would be suspicious; the bracket is a 95% block-bootstrap range.
**Brier skill** > 0 means the probabilities beat just quoting the normal odds. **Log loss** punishes confident mistakes (lower is better). **Rank IC** = rank correlation between the probability and the actual outcome.
""")
        rows = []
        for name, wf in ml.items():
            wk = known(wf, fwd, unseen_idx)
            if wk.empty:
                continue
            yy = (fwd.loc[wk.index] > 0).astype(int)
            auc = roc_auc_score(yy, wk["p"]) if yy.nunique() > 1 else np.nan
            lo_a, hi_a = auc_ci(yy.values, wk["p"].values, h)
            br, bss, ll = prob_metrics(wk, yy)
            sm = summarize(fwd, ml_verdict(wk, margin), wk.index, h).set_index("Group")
            ok = sm.loc[FAV, "Months"] >= 3 and sm.loc[UNF, "Months"] >= 3
            sp_ = (sm.loc[FAV, "Avg"] - sm.loc[UNF, "Avg"]) * 100 if ok else np.nan
            rows.append({"Model": name, "AUC": f"{auc:.2f} [{lo_a:.2f}–{hi_a:.2f}]" if pd.notna(lo_a) else f"{auc:.2f}",
                         "Brier skill": f"{bss:+.3f}", "Log loss": f"{ll:.3f}", "Rank IC": f"{rank_ic(wk['p'], fwd.loc[wk.index]):+.2f}",
                         "Accuracy": f"{((wk['p'] > 0.5) == yy).mean():.0%}", "Naive": f"{((wk['base'] > 0.5) == yy).mean():.0%}",
                         "Favorable minus Unfavorable": "-" if np.isnan(sp_) else f"{sp_:+.1f} pts"})
        st.subheader(f"Out-of-sample comparison ({unseen_lbl} months)")
        show_df(st, pd.DataFrame(rows))
        st.caption("Simple vs complex: if logistic regression scores about the same as the forests, the extra complexity is not buying anything.")
        wf = ml[ml_choice]
        wk = known(wf, fwd, unseen_idx)
        if wk.empty:
            st.warning("Not enough history to train. Move the start date earlier.")
        else:
            yy = (fwd.loc[wk.index] > 0).astype(int)
            auc = roc_auc_score(yy, wk["p"]) if yy.nunique() > 1 else np.nan
            lo_a, hi_a = auc_ci(yy.values, wk["p"].values, h)
            _, bss, _ = prob_metrics(wk, yy)
            skill = pd.notna(lo_a) and lo_a > 0.5 and bss > 0
            (st.success if skill else st.warning)(
                f"{ml_choice}: AUC {auc:.2f}" + (f" (95% range {lo_a:.2f}–{hi_a:.2f})" if pd.notna(lo_a) else "") + f", Brier skill {bss:+.3f}. "
                + ("The AUC range sits above 0.50, so there is some skill, but the test period is short." if skill
                   else "The AUC range includes 0.50 or the probabilities do not beat the normal odds: little reliable skill on unseen data. Don't rely on it."))
            msum = summarize(fwd, ml_verdict(wk, margin), wk.index, h, fdd)
            st.subheader(f"{ml_choice}: outcome after each call")
            show_df(st, fmt_summary(msum, True))
            c1, c2 = st.columns(2)
            pf = go.Figure(go.Scatter(x=wf.index, y=wf["p"], name="P(positive)"))
            pf.add_scatter(x=wf.index, y=wf["base"], name="Normal odds", line=dict(dash="dash"))
            pf.update_layout(title="Out-of-sample probability", yaxis_tickformat=".0%", height=340)
            show_plot(c1, pf)
            try:
                bins = pd.qcut(wk["p"], 5, duplicates="drop")
                cal = pd.DataFrame({"pred": wk["p"].groupby(bins, observed=True).mean(), "act": yy.groupby(bins, observed=True).mean()})
                cf = go.Figure(go.Scatter(x=cal["pred"], y=cal["act"], mode="lines+markers", name="Model"))
                cf.add_scatter(x=[0, 1], y=[0, 1], name="Perfect", line=dict(dash="dash"))
                cf.update_layout(title="Calibration: predicted vs actual", xaxis_title="Predicted", yaxis_title="Actual", height=340)
                show_plot(c2, cf)
            except ValueError:
                c2.info("Not enough spread in predictions for a calibration plot.")
            if ml_contrib is not None:
                st.subheader("Why today's prediction? (log-odds contribution of each indicator)")
                st.bar_chart(ml_contrib.rename(index={c: META[c][0] for c in ml_contrib.index}).sort_values())
                st.caption("Positive pushes toward the target being positive, negative pushes away. Exact for logistic regression.")
            else:
                st.subheader("What the model uses most")
                st.bar_chart(ml_imp.rename(index={c: META[c][0] for c in ml_imp.index}).sort_values())
                st.caption("Importance is not causality, and correlated indicators share credit unpredictably. Switch to logistic regression for exact contributions.")

# ------------------------------------------------------------------ scan
if auto_h:
    with T["scan"]:
        st.markdown(f"""
Each look-ahead is fitted on the **first 70% of the learning period** and scored on its **last 30%**. **IC** = rank correlation between prediction and outcome (0 = no skill, ~0.10 = decent).
**Stable IC** is the worse of the two validation halves, so a horizon only wins if it works in both. Selected: **{h} months**, then locked. Validation and final test are never used here.
""")
        d = scan.copy()
        d["Look-ahead"] = d["h"].map(lambda x: f"{x} months" + (" ✅" if x == h else ""))
        for c in ["rules_ic", "ml_ic", "combined", "stable"]:
            d[c] = d[c].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}")
        d["Independent obs."] = (scan["n_val"] / scan["h"]).round(0).astype(int)
        show_df(st, d[["Look-ahead", "rules_ic", "ml_ic", "combined", "stable", "Independent obs."]].rename(
            columns={"rules_ic": "Rules IC", "ml_ic": "ML IC", "combined": "Average IC", "stable": "Stable IC"}))
        st.caption("Longer look-aheads have far fewer independent observations. If every IC is near zero, there is no reliable horizon, and that is a valid answer.")

# ------------------------------------------------------------------ bitcoin & gold now
with T["now"]:
    st.subheader(f"✅ Bitcoin & gold: where the environment stands now ({now:%b %Y})")
    st.caption(f"Each market gets its own rules and ML read for the next {h} months, using the same indicators and settings. "
               "Favorable / Unfavorable describes how similar environments played out historically. It is not a buy or sell instruction. "
               "These three reads always use plain return targets (bitcoin, gold, bitcoin minus gold), whatever target is selected in the sidebar.")

    for col, k in zip(st.columns(3), NOW_KINDS):
        r = NOW[k]
        box = card(col)
        box.markdown(f"### {r['icon']} {NOW_NAMES[k]}")
        box.markdown(f"**{REL_TXT[r['label']] if k == 'gold' else r['label']}**")
        box.markdown(f"**Confidence: {r['conf']}** ({r['agree']} of {r['n_votes']} evidence channels agree)")
        box.markdown(" · ".join(f"{n_} {VICON[v]}" for n_, v in r["votes"].items()))
        box.caption("Macro models = " + ", ".join(f"{n_} {VICON[v]}" for n_, v in r["comps"].items()))
        box.markdown(f"Rules: {ICON[r['rules']]} {r['rules']} (score {r['score']:+.1f})")
        if r["ml"]:
            box.markdown(f"ML: {r['ml']['p']:.0%} chance {NOW_PHRASE[k]} in {h}m (normal {r['ml']['base']:.0%}) → {ICON[r['ml']['verdict']]} {r['ml']['verdict']}")
        box.markdown(f"Similar past periods (descriptive): median {pc(r['ana_med'])}, {r['ana_pos']:.0%} positive")
        box.markdown(f"Evidence on unseen data: **{r['ev_level']}**")
        box.caption(r["ev_text"])
        if r["ml"] and pd.notna(r["ml"]["auc"]):
            box.caption(f"ML on unseen data: AUC {r['ml']['auc']:.2f}, Brier skill {r['ml']['bss']:+.3f}.")
        if r["why_good"]:
            box.markdown("**Helping now:** " + "; ".join(r["why_good"][:4]))
        if r["why_bad"]:
            box.markdown("**Hurting now:** " + "; ".join(r["why_bad"][:4]))

    es, eg, er = NOW["ret"]["e"], NOW["goldabs"]["e"], NOW["gold"]["e"]
    sl, gl = NOW["ret"]["label"].lower(), NOW["goldabs"]["label"].lower()
    if es >= 0.7 and eg >= 0.7:
        msg = "Both bitcoin and gold sit in historically favorable environments."
    elif es <= 0.3 and eg <= 0.3:
        msg = "Both bitcoin and gold sit in historically unfavorable environments."
    elif es >= 0.7:
        msg = f"Bitcoin's environment is historically favorable ({sl}), while gold's is {gl}."
    elif eg >= 0.7:
        msg = f"Gold's environment is historically favorable ({gl}), while bitcoin's is {sl}."
    elif es <= 0.3:
        msg = f"Bitcoin's environment is historically unfavorable ({sl}), while gold's is {gl}."
    elif eg <= 0.3:
        msg = f"Gold's environment is historically unfavorable ({gl}), while bitcoin's is {sl}."
    else:
        msg = "Neither market is in a clearly favorable or unfavorable environment: mostly neutral."
    if er >= 0.7:
        msg += " Similar conditions have historically favored bitcoin over gold."
    elif er <= 0.3:
        msg += " Similar conditions have historically favored gold over bitcoin."
    else:
        msg += " There is no clear bitcoin-versus-gold preference."
    (st.success if max(es, eg) >= 0.7 else (st.warning if min(es, eg) <= 0.3 else st.info))(msg)

    weak = [NOW_NAMES[k] for k in NOW_KINDS if NOW[k]["ev_level"] != "Strong"]
    if weak:
        st.caption("Evidence is Weak, None or Unknown for: " + ", ".join(weak) + ". For those, the label describes today's conditions "
                   "but the rules did not reliably predict outcomes on data they had never seen, so treat it as context, not a signal.")

    show_df(st, pd.DataFrame([{
        "Market": NOW_NAMES[k],
        "Overall": f"{NOW[k]['icon']} {REL_TXT[NOW[k]['label']] if k == 'gold' else NOW[k]['label']}",
        "Confidence": f"{NOW[k]['conf']} ({NOW[k]['agree']}/{NOW[k]['n_votes']} agree)",
        "Rules": f"{ICON[NOW[k]['rules']]} {NOW[k]['rules']}",
        "ML chance": "-" if not NOW[k]["ml"] else f"{NOW[k]['ml']['p']:.0%} (normal {NOW[k]['ml']['base']:.0%})",
        "Similar periods, median": pc(NOW[k]["ana_med"]),
        "Evidence": NOW[k]["ev_level"]} for k in NOW_KINDS]))
    st.caption("Confidence combines two things: whether the two evidence channels (macro models, which share the same indicators, and price trend) point the same way, "
               "and whether the rules beat chance on unseen data. It is a rough guide, not a probability.")
    st.warning("**Educational only, not financial advice.** This reads macro conditions only. It does not know about valuation, news, taxes, "
               "your goals or your time horizon, and relationships that held in the past can stop working.")


# ================================================================== Word report tab (last tab)
# Collects the results of every tab into one Word file. Built only when the button is pressed (not on every rerun).
# Needs: python-docx (required), kaleido (optional, only for the charts).
import io
import re

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    _HAS_DOCX = True
except Exception:  # noqa: BLE001
    _HAS_DOCX = False

_RP_REPL = {"✅": "Yes", "❌": "No", "✔": "Yes", "✘": "No", "🟢": "(+)", "🔴": "(-)", "⚪": "(0)", "🟡": "(~)"}
_RP_EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F\u200d]")
_RP_IMG_OK = [True]
_RP_VOTE = {1: "favorable", 0: "neutral", -1: "unfavorable"}


def _rp_clean(x):
    s = "" if x is None else str(x)
    for a, b in _RP_REPL.items():
        s = s.replace(a, b)
    return re.sub(r"[ \t]{2,}", " ", _RP_EMOJI.sub("", s)).strip()


def _rp_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _rp_field(par, instr, size=8):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(size)
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        elif kind == "text":
            r.text = txt
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def _rp_widths(df, total=17.4, minw=1.2):
    """Column widths (cm) proportional to the longest text in each column, summing to the page width."""
    w = []
    for c in df.columns:
        longest = max([len(w_) for w_ in _rp_clean(c).split()] + [0])
        cells = [min(len(_rp_clean(v)), 45) for v in df[c].head(60)]
        if max(cells + [0]) <= 3 and longest <= 5:  # rank / index columns
            w.append(3.5)
        else:
            w.append(max(7.5, longest * 1.1, (sum(cells) / max(len(cells), 1)) * 0.8, max(cells + [0]) * 0.45))
    tot = sum(w)
    ws = [max(minw, total * x / tot) for x in w]
    k = total / sum(ws)
    return [x * k for x in ws]


class _Rep:
    """Thin wrapper around python-docx with the handful of building blocks the report needs."""

    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
        nrm = self.d.styles["Normal"]
        nrm.font.name, nrm.font.size = "Calibri", Pt(10)
        nrm.paragraph_format.space_after = Pt(4)
        for name, size, color in (("Heading 1", 17, "1F3864"), ("Heading 2", 13, "2F5496"), ("Heading 3", 11, "2F5496")):
            s = self.d.styles[name]
            s.font.name, s.font.size, s.font.bold = "Calibri", Pt(size), True
            s.font.color.rgb = RGBColor.from_string(color)
            rf = s.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            s.paragraph_format.keep_with_next = True
        self.d.styles["Heading 1"].paragraph_format.page_break_before = True
        self.d.styles["Heading 1"].paragraph_format.space_after = Pt(8)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Bitcoin Macro Environment Analyzer · educational only, not financial advice · page ")
        r.font.size = Pt(8)
        _rp_field(fp, "PAGE")

    # -- text
    def _runs(self, par, text, size=None, color=None, italic=False, bold=False):
        for i, part in enumerate(_rp_clean(text).split("**")):
            if not part:
                continue
            r = par.add_run(part)
            r.bold, r.italic = bold or i % 2 == 1, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def h1(self, t):
        self.d.add_heading(_rp_clean(t), 1)

    def h2(self, t):
        self.d.add_heading(_rp_clean(t), 2)

    def h3(self, t):
        self.d.add_heading(_rp_clean(t), 3)

    def p(self, t, **kw):
        par = self.d.add_paragraph()
        self._runs(par, t, **kw)
        return par

    def note(self, t):
        return self.p(t, size=8.5, color="595959", italic=True)

    def bullets(self, items):
        for it in items:
            self._runs(self.d.add_paragraph(style="List Bullet"), it)

    def box(self, t, kind="info"):
        fill = {"ok": "E2F0D9", "warn": "FFF2CC", "info": "DDEBF7", "bad": "F8D7DA"}[kind]
        tb = self.d.add_table(rows=1, cols=1)
        tb.style = "Table Grid"
        c = tb.rows[0].cells[0]
        c.width = Cm(17.4)
        _rp_shade(c, fill)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        self._runs(c.paragraphs[0], t, size=9.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- tables and charts
    def table(self, df, font=None, widths=None):
        if df is None or len(df) == 0:
            self.note("No data for this table.")
            return
        df = df.reset_index(drop=True)
        n = len(df.columns)
        font = font or (8 if n <= 6 else (7.5 if n <= 8 else 7))
        tb = self.d.add_table(rows=1, cols=n)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(df.columns):
            cell = tb.rows[0].cells[i]
            _rp_shade(cell, "1F3864")
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            self._runs(par, str(c), size=font, color="FFFFFF", bold=True)
        trPr = tb.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        for ri, row in enumerate(df.itertuples(index=False)):
            cells = tb.add_row().cells
            cs = OxmlElement("w:cantSplit")
            tb.rows[-1]._tr.get_or_add_trPr().append(cs)
            for i, v in enumerate(row):
                if ri % 2:
                    _rp_shade(cells[i], "F2F2F2")
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                self._runs(par, "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v), size=font)
        widths = widths or _rp_widths(df)
        tb.autofit = False
        tblPr = tb._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for i, w in enumerate(widths):
            tb.columns[i].width = Cm(w)
        for row in tb.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def kv(self, pairs, font=8.5):
        self.table(pd.DataFrame(pairs, columns=["Item", "Value"]), font=font, widths=[5.5, 11.9])

    def image(self, fig, title=None, w=17.2, h_px=520):
        if not _RP_IMG_OK[0] or fig is None:
            return False
        try:
            f2 = go.Figure(fig)
            f2.update_layout(template="plotly_white", paper_bgcolor="white", plot_bgcolor="white", title=title or (f2.layout.title.text if f2.layout.title else None))
            png = f2.to_image(format="png", width=1100, height=h_px, scale=2)
            self.d.add_picture(io.BytesIO(png), width=Cm(w))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
        except Exception:  # noqa: BLE001  (kaleido missing or failed: the report simply has no charts)
            _RP_IMG_OK[0] = False
            return False

    def save(self):
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


def _rp_v(name, default=None):
    return globals().get(name, default)


def _rp_ml_stats(name):
    wk = known(ml[name], fwd, unseen_idx)
    if wk.empty:
        return None
    yy = (fwd.loc[wk.index] > 0).astype(int)
    auc = roc_auc_score(yy, wk["p"]) if yy.nunique() > 1 else np.nan
    lo_a, hi_a = auc_ci(yy.values, wk["p"].values, h)
    _, bss, ll = prob_metrics(wk, yy)
    sm = summarize(fwd, ml_verdict(wk, margin), wk.index, h).set_index("Group")
    ok = sm.loc[FAV, "Months"] >= 3 and sm.loc[UNF, "Months"] >= 3
    return dict(wk=wk, yy=yy, auc=auc, lo=lo_a, hi=hi_a, bss=bss, ll=ll, ic=rank_ic(wk["p"], fwd.loc[wk.index]),
                acc=float(((wk["p"] > 0.5) == yy).mean()), naive=float(((wk["base"] > 0.5) == yy).mean()),
                spr=(sm.loc[FAV, "Avg"] - sm.loc[UNF, "Avg"]) * 100 if ok else np.nan)


def _rp_gap_text(tb, label):
    """Favorable-minus-Unfavorable sentence from a summarize(...).set_index('Group') frame, or None if too few months."""
    if tb.loc[FAV, "Months"] >= 3 and tb.loc[UNF, "Months"] >= 3:
        sp_ = (tb.loc[FAV, "Avg"] - tb.loc[UNF, "Avg"]) * 100
        ov = bool(tb.loc[FAV, "lo"] <= tb.loc[UNF, "hi"])
        return (f"{label}: Favorable minus Unfavorable = {sp_:.1f} points. "
                + ("The confidence ranges overlap, so this could easily be noise." if ov else "The confidence ranges do not overlap."),
                "ok" if sp_ > 2 and not ov else "warn")
    return None


# ---------------------------------------------------------------- sections
def _rp_summary(R):
    R.p(f"This report collects the results of every tab of the Bitcoin Macro Environment Analyzer, using data through **{m.index[-1]:%B %Y}** "
        f"(latest month in the models: **{now:%B %Y}**). It describes how bitcoin and gold behaved after macro conditions similar to today's. "
        "It is **not** a buy or sell recommendation.")
    R.h2("Bitcoin, gold and bitcoin versus gold now")
    t = []
    for k in NOW_KINDS:
        r = NOW[k]
        t.append({"Market": NOW_NAMES[k], "Overall read": REL_TXT[r["label"]] if k == "gold" else r["label"],
                  "Confidence": f"{r['conf']} ({r['agree']} of {r['n_votes']} channels agree)", "Rules": f"{r['rules']} (score {r['score']:+.1f})",
                  "ML chance": "-" if not r["ml"] else f"{r['ml']['p']:.0%} (normal {r['ml']['base']:.0%})", "Evidence on unseen data": r["ev_level"]})
    R.table(pd.DataFrame(t), font=8)
    es, eg, er = NOW["ret"]["e"], NOW["goldabs"]["e"], NOW["gold"]["e"]
    sl, gl = NOW["ret"]["label"].lower(), NOW["goldabs"]["label"].lower()
    if es >= 0.7 and eg >= 0.7:
        msg = "Both bitcoin and gold sit in historically favorable environments."
    elif es <= 0.3 and eg <= 0.3:
        msg = "Both bitcoin and gold sit in historically unfavorable environments."
    elif es >= 0.7:
        msg = f"Bitcoin's environment is historically favorable ({sl}), while gold's is {gl}."
    elif eg >= 0.7:
        msg = f"Gold's environment is historically favorable ({gl}), while bitcoin's is {sl}."
    elif es <= 0.3:
        msg = f"Bitcoin's environment is historically unfavorable ({sl}), while gold's is {gl}."
    elif eg <= 0.3:
        msg = f"Gold's environment is historically unfavorable ({gl}), while bitcoin's is {sl}."
    else:
        msg = "Neither market is in a clearly favorable or unfavorable environment: mostly neutral."
    msg += (" Similar conditions have historically favored bitcoin over gold." if er >= 0.7 else
            (" Similar conditions have historically favored gold over bitcoin." if er <= 0.3 else " There is no clear bitcoin-versus-gold preference."))
    R.box(msg, "ok" if max(es, eg) >= 0.7 else ("warn" if min(es, eg) <= 0.3 else "info"))

    R.h2("Key numbers")
    items = [f"**Rules verdict ({now:%b %Y}):** {v_now} (macro score {sc_now:+.1f})."]
    if run_ml:
        items.append(f"**Machine learning ({ml_choice}):** {ml_p:.0%} chance that {tgt_txt} in {h} months (normal odds {ml_base:.0%}).")
    items.append(f"**Suggested exposure:** {expo_now:.0%} of the bitcoin allocation you already intended (Favorable 100%, Neutral 50%, Unfavorable 0%).")
    cur, n_cur, med, n_h, avg_h = streak
    items.append(f"**Persistence:** {cur.lower()} for {n_cur} month(s) (typical run {med:.0f}).")
    if reg is not None:
        items.append(f"**Macro regime:** {reg.iloc[-1]}.")
    a_ = ana["After"]
    items.append(f"**10 most similar past environments (descriptive):** median outcome {pc(a_.median())}, {(a_ > 0).mean():.0%} positive after {h} months.")
    if PROJ_REPORT.get("head"):
        hd, L_ = PROJ_REPORT["head"], PROJ_REPORT["L"]
        items.append(f"**Projection ({L_} months):** median ${hd['median']:,.1f} ({hd['chg']:+.1%} vs ${PROJ_REPORT['S0']:,.1f} now), 80% range ${hd['lo']:,.0f} to ${hd['hi']:,.0f}, "
                     f"{hd['up']:.0%} chance of being higher.")
    xc = _rp_v("_XC", [])
    if "ex_recess" in xc and F_full["ex_recess"].notna().any():
        items.append(f"**Recession probability, next 12 months (yield-curve model):** {F_full['ex_recess'].dropna().iloc[-1]:.0f}%.")
    if "ex_ff_chg" in xc and F_full["ex_ff_chg"].notna().any():
        items.append(f"**Fed funds priced over 12 months:** {F_full['ex_ff_chg'].dropna().iloc[-1]:+.2f} points.")
    R.bullets(items)

    R.h2("How much can this be trusted?")
    rel = []
    tb = summarize(fwd, verdict, grade_idx, h, fdd).set_index("Group")
    g_ = _rp_gap_text(tb, f"Rules, {grade_name.lower()} period")
    rel.append(g_[0] if g_ else f"Rules, {grade_name.lower()} period: too few Favorable or Unfavorable months to compare.")
    if run_ml:
        s_ = _rp_ml_stats(ml_choice)
        if s_:
            rel.append(f"{ml_choice} on unseen data: AUC {s_['auc']:.2f}"
                       + (f" (95% range {s_['lo']:.2f} to {s_['hi']:.2f})" if pd.notna(s_["lo"]) else "") + f", Brier skill {s_['bss']:+.3f}.")
    if PROJ_REPORT.get("n_hor"):
        rel.append(f"Projection fan: beat history-only on {PROJ_REPORT['n_good']} of {PROJ_REPORT['n_hor']} horizons (average skill {PROJ_REPORT['avg_skill']:+.1%}); "
                   f"the 80% range held the outcome {PROJ_REPORT['cov']:.0%} of the time.")
    weak = [NOW_NAMES[k] for k in NOW_KINDS if NOW[k]["ev_level"] != "Strong"]
    if weak:
        rel.append("Evidence on unseen data is Weak, None or Unknown for: " + ", ".join(weak) + ". For those, the label describes today's conditions but is context, not a signal.")
    R.bullets(rel)
    R.note("Contents: Dashboard · Guide · Environments · Indicator ranking · Out-of-sample · Regimes & analogues · Backtest · Robustness · Explorer · "
           "Expectations · Projection · Machine learning · Look-ahead scan · Bitcoin & gold now.")


def _rp_settings(R):
    per = [("Learn period", f"{learn_idx[0]:%b %Y} to {learn_idx[-1]:%b %Y}"), ("Validation period", f"{val_idx[0]:%b %Y} to {val_idx[-1]:%b %Y}"),
           ("Final test period", f"{test_idx[0]:%b %Y} to {test_idx[-1]:%b %Y}" if REVEAL else "hidden (research mode)")]
    R.kv([("Data through", f"{m.index[-1]:%b %Y}"), ("Data start date", start), ("Evaluation mode", eval_mode), ("What is predicted", target_name),
          ("Look-ahead", f"{h} months" + (" (auto-selected on the learn period)" if auto_h else "")), *per,
          ("Indicators / factors in use", len(META)), ("Factor compression", f"on (merge at |corr| ≥ {fac_thr:.2f})" if fac_mode else "off"),
          ("Expectation indicators in every model", "yes" if use_exp else "no (still shown in the Expectations and Projection tabs)"),
          ("First-release (ALFRED) data", "yes" if use_vintage else "no"), ("Low / Mid / High defined by", scheme),
          ("Environments selected by", sel_mode + (f" (|t| ≥ {t_thr:.1f})" if not sel_mode.startswith("FDR") else f" (q ≤ {q_thr:.2f})")),
          ("Indicator weighting", wlabel), ("Machine learning", f"{ml_choice} in detail, retrained every {step} months, margin {margin:.2f}" if run_ml else "off"),
          ("Backtest strategy", mode), ("Implementation", f"{impl_key}, {bps} bps per 100% traded"),
          ("Projection", f"{proj_model}; features: {proj_feats}; {proj_len} months; trust {proj_trust}%")])
    R.p("**Indicators in use:** " + "; ".join(META[c][0] for c in META) + ".", size=9)
    if failed:
        R.note("Data series that failed to download and were skipped: " + ", ".join(failed) + ".")
    if REVEAL and peek:
        R.box(f"Final test revealed. It has been shown under {peek[0]} different set(s) of settings"
              + (" (this one is new)." if peek[1] else " (this one was already counted).")
              + " Only the first reveal is a clean test; every extra look makes the final test a selection tool, so treat results as optimistic.",
              "ok" if peek[0] == 1 else "warn")
    elif not REVEAL:
        R.box("Research mode: the final test period is hidden everywhere in this report and all grades use the validation period.", "info")


def _rp_dash(R):
    R.p(f"**Historically {v_now.lower()} environment for bitcoin ({now:%b %Y}).**")
    k = [("Macro score (rules)", f"{sc_now:+.1f}" if cfg[3] != "equal" else f"{sc_now:+.0f}")]
    if run_ml:
        k.append((f"ML: chance {tgt_txt} in {h} months", f"{ml_p:.0%} ({(ml_p - ml_base) * 100:+.1f} pts vs normal {ml_base:.0%})"))
    k.append(("Suggested exposure", f"{expo_now:.0%} of the intended bitcoin allocation"))
    if "mvrv" in F_all and pd.notna(F_all["mvrv"].get(now, np.nan)):
        k.append(("MVRV (market cap / realized cap)", f"{F_all['mvrv'][now]:.2f} (on-chain valuation gauge, not a timing signal)"))
    R.kv(k)
    r_ = NOW["ret"]
    R.p(f"**Bitcoin overall:** {r_['label']} | confidence **{r_['conf']}** ({r_['agree']} of {r_['n_votes']} evidence channels agree; evidence on unseen data: {r_['ev_level']}). "
        + " · ".join(f"{n_}: {_RP_VOTE[v]}" for n_, v in r_["votes"].items())
        + ". Macro models = " + ", ".join(f"{n_}: {_RP_VOTE[v]}" for n_, v in r_["comps"].items()) + " (all reading the same indicators).")

    R.h2("Pillar scores (-100 = all unfavorable, +100 = all favorable)")
    S_now, rows_, psc = S.loc[now], [], {}
    for pl in PILLARS:
        cols = [c for c in META if META[c][4] == pl]
        if cols:
            sc_ = 100 * (sum((c, S_now[c]) in good for c in cols) - sum((c, S_now[c]) in bad for c in cols)) / len(cols)
            psc[pl] = sc_
            rows_.append({"Pillar": pl, "Score": f"{sc_:+.0f}", "Read": "Favorable" if sc_ > 15 else ("Unfavorable" if sc_ < -15 else "Mixed"), "Indicators": len(cols)})
    R.table(pd.DataFrame(rows_))
    mon, ind = [psc[p] for p in MONETARY_PILLARS if p in psc], [psc[p] for p in INDUSTRIAL_PILLARS if p in psc]
    if mon and ind:
        R.p(f"**Bitcoin's two backdrops:** macro and liquidity backdrop (rates, dollar, M2, Fed liquidity, growth, risk appetite) {np.mean(mon):+.0f}; on-chain backdrop (valuation, miners, network activity) {np.mean(ind):+.0f}.")
    if reg is not None:
        R.p(f"**Macro regime:** {reg.iloc[-1]}")
    cur, n_cur, med, n_h, avg_h = streak
    R.p(f"**Persistence:** {cur.lower()} for {n_cur} month(s) (typical run {med:.0f}). "
        + (f"After 3+ months in this state, the average outcome was {avg_h:+.1%} ({n_h} months)." if n_h >= 5 else ""))

    R.h2("Where each indicator stands today")
    st_rows = []
    for c in META:
        s_ = S_now[c]
        st_rows.append({"Indicator": META[c][0], "Pillar": META[c][4], "Current value": fmt_val(c, F_ok.loc[now, c]), "State": state_label(c, s_),
                        "Effect on the verdict": "helps" if (c, s_) in good else ("hurts" if (c, s_) in bad else "no active rule")})
    R.table(pd.DataFrame(st_rows))

    R.h2("10 most similar past environments (descriptive, not tradable)")
    a_ = ana["After"]
    R.p(f"Outcome after {h} months ({target_name.lower()}):")
    R.kv([("Median", pc(a_.median())), ("Average", pc(a_.mean())), ("Positive", f"{(a_ > 0).mean():.0%}"), ("Worst", pc(a_.min())), ("Best", pc(a_.max()))])
    R.note("Only 10 observations: indicative, not statistical proof. The search uses all available history, so it is not a trading signal.")

    th = state_thresholds(F_ok, learn_idx, pit)
    rows_ = []
    for dct, eff in ((good, "helps"), (bad, "hurts")):
        for c, s_ in sorted(dct):
            if S_now[c] != s_:
                continue
            lo_, hi_ = th.loc[c].iloc[0], th.loc[c].iloc[1]
            flips = {"Low": f"rises above {fmt_val(c, lo_)}", "High": f"falls below {fmt_val(c, hi_)}", "Mid": f"leaves {fmt_val(c, lo_)} to {fmt_val(c, hi_)}"}[s_]
            rows_.append({"Indicator": META[c][0], "Now": state_label(c, s_), "Effect": eff, "Current value": fmt_val(c, F_ok.loc[now, c]), "Stops counting if it": flips})
    R.h2("What would change the rules' verdict?")
    if rows_:
        R.table(pd.DataFrame(rows_))
    else:
        R.p("No active rule is currently contributing to the verdict.")


def _rp_guide(R):
    R.p(f"**Three periods.** History is split into LEARN ({learn_idx[0]:%b %Y} to {learn_idx[-1]:%b %Y}), VALIDATION ({val_idx[0]:%b %Y} to {val_idx[-1]:%b %Y}) and FINAL TEST "
        f"({test_idx[0]:%b %Y} to {test_idx[-1]:%b %Y}). Indicators, factors, rules and the look-ahead are chosen on LEARN only; strategy and signal on VALIDATION; "
        "the FINAL TEST only grades the locked choice. Outcome windows are purged so periods never overlap.")
    R.bullets([
        f"**What is predicted:** {target_name}, {h} months ahead. Drawdown targets predict risk instead of direction; relative targets ask whether bitcoin beats gold or cash.",
        f"**Environments (rules):** each month is described by {len(META)} {'factors / indicators' if fac_mode else 'macro indicators'}, each split into Low / Neutral / High. "
        "Environments that are statistically clear and hold in both halves of the learn period become rules.",
        "**Machine learning:** models estimate the chance of a positive outcome, retrained walk-forward on data whose outcome was already known at the time.",
        "**Analogues:** the 10-closest-months table is descriptive and not tradable; the predictive analogue test searches only months whose outcome was known then.",
        "**Backtest:** strategies on ETFs or futures with trading costs, compared with a no-timing mix that has the same average exposure, plus a randomization test.",
        "**Expectations:** eleven forward-looking indicators (market-implied, published model, or walk-forward model-implied), graded against 'nothing changes'.",
        "**Evidence channels:** rules, ML and analogues read the same indicators, so they count as one macro channel; price trend is the second.",
    ])
    R.h2("Known limits")
    R.bullets(["Publication lags are approximations wherever first-release data is off or unavailable.",
               "Spot / ETF-style uses BTC-USD and GLD prices with no fee deducted; futures-style adds T-bill collateral income but does not model the CME basis.",
               "On-chain data comes from free sources that can be revised; exchange flows, ETF flows and funding rates are not included. Bitcoin has only about three halving cycles of history.",
               "Selection bias remains whenever many settings or targets are tried. Judge settings by the final test and the randomization results, not by the best-looking one.",
               "Educational only, not financial advice."])


def _rp_env(R):
    R.p(f"Average outcome after {h} months in each Low / Neutral / High environment (learn period). {len(stats)} environments were tested; "
        f"{len(good)} became favorable rules and {len(bad)} unfavorable rules.")
    rows_ = []
    for c in META:
        row = {"Indicator": META[c][0]}
        for s_ in STATES:
            r = stats[(stats["col"] == c) & (stats["state"] == s_)]
            row[s_] = "-" if r.empty else f"{state_label(c, s_)}: {r['avg'].iloc[0]:+.1%} (n={int(r['n'].iloc[0])})"
        rows_.append(row)
    R.table(pd.DataFrame(rows_), font=7)

    def rank_table(df):
        return pd.DataFrame({"Environment": [state_label(r.col, r.state) for r in df.itertuples()], "Indicator": [META[r.col][0] for r in df.itertuples()],
                             "Months": df["n"].values, "Avg after": [pc(v) for v in df["avg"]], "vs. average": [f"{v * 100:+.1f} pts" for v in df["edge"]],
                             "% positive": [pc(v, False) for v in df["win"]], "HAC t": [f"{v:+.1f}" for v in df["t"]], "FDR q": [f"{v:.2f}" for v in df["q"]],
                             "Both halves?": ["Yes" if v else "No" for v in df["consistent"]]})

    R.h2("Most favorable environments")
    R.table(rank_table(stats.sort_values("edge", ascending=False).head(6)))
    R.h2("Most unfavorable environments")
    R.table(rank_table(stats.sort_values("edge").head(6)))
    R.note("HAC t above ±2 is fairly strong. Because many environments are tested, some look good by luck: the FDR q column adjusts for that.")
    R.h2("Rules in use")
    R.p("**Favorable:** " + ("; ".join(state_label(c, s_) for c, s_ in sorted(good)) if good else "none"))
    R.p("**Unfavorable:** " + ("; ".join(state_label(c, s_) for c, s_ in sorted(bad)) if bad else "none"))


def _rp_rank(R):
    rk_ = _rp_v("rk")
    if rk_ is None:
        rk_ = indicator_ranking(F_ok, fwd, S, learn_idx, unseen_idx, h, fdd)
    ab_ = _rp_v("ab")
    if ab_ is None:
        ab_ = ablation(F_ok, fwd, h, sp["cut"], step, unseen_idx.values).set_index("col")
    R.p(f"Target: **{target_name}** over **{h} months**. Each indicator is compared with the outcome that followed, first on the learn period, then on unseen data "
        f"({unseen_idx[0]:%b %Y} to {unseen_idx[-1]:%b %Y}, {unseen_lbl}). Ranking is by the smaller of the two correlations, and 0 if the direction flips.")
    top_ = rk_[rk_["minic"] > 0].head(3)
    if len(top_):
        R.box("Most consistent so far: " + ", ".join(f"{META[r.col][0]} ({r.minic:.0f}%)" for r in top_.itertuples()), "ok")
    else:
        R.box("No indicator kept the same direction on unseen data. Treat every single indicator as unreliable for this target and look-ahead.", "warn")
    R.table(pd.DataFrame({
        "#": range(1, len(rk_) + 1), "Indicator": [META[c][0] for c in rk_["col"]],
        "Direction": ["-" if pd.isna(v) else ("Higher → better" if v > 0 else "Higher → worse") for v in rk_["ic_l"]],
        "Learn corr": [pc(v) for v in rk_["ic_l"]], "Unseen corr": [pc(v) for v in rk_["ic_t"]], "Same dir.?": ["Yes" if v else "No" for v in rk_["same"]],
        "Min |corr|": [f"{v:.0f}%" for v in rk_["minic"]], "HAC t": ["-" if pd.isna(v) else f"{v:+.1f}" for v in rk_["t_l"]],
        "Hit rate (unseen)": ["-" if pd.isna(a) else f"{a:.0%} ± {b * 100:.0f}" for a, b in zip(rk_["hit"], rk_["hit_ci"])],
        "ΔAUC if removed": ["-" if pd.isna(ab_["dauc"].get(c, np.nan)) else f"{ab_['dauc'][c]:+.3f}" for c in rk_["col"]],
        "Verdict": rk_["status"]}))
    R.h2("Risk-aware view (descriptive)")
    R.table(pd.DataFrame({
        "Indicator": [META[c][0] for c in rk_["col"]],
        "High − Low, learn": ["-" if pd.isna(v) else f"{v:+.1f} pts" for v in rk_["hl_l"]], "High − Low, unseen": ["-" if pd.isna(v) else f"{v:+.1f} pts" for v in rk_["hl_t"]],
        "In σ, learn": ["-" if pd.isna(v) else f"{v:+.2f}σ" for v in rk_["eff_l"]], "In σ, unseen": ["-" if pd.isna(v) else f"{v:+.2f}σ" for v in rk_["eff_t"]],
        "Drawdown after High − Low (unseen)": ["-" if pd.isna(v) else f"{v:+.0f} pts" for v in rk_["dd_t"]]}))
    R.note("Corr = rank correlation with the outcome (±10% is already useful for macro data, below ±5% is hard to tell from noise). Min |corr| is a consistency score, not a test. "
           "ΔAUC if removed: positive = the indicator adds information (differences under about 0.02 are noise). With this many indicators, a few look good by luck.")
    pill = rk_.assign(Pillar=[META[c][4] for c in rk_["col"]]).groupby("Pillar")["minic"].agg(["mean", "max", "count"]).sort_values("mean", ascending=False)
    R.h2("Which themes carry the most consistent signal")
    R.table(pd.DataFrame({"Pillar": pill.index, "Average Min |corr|": [f"{v:.1f}%" for v in pill["mean"]], "Best indicator in pillar": [f"{v:.1f}%" for v in pill["max"]],
                          "Indicators": pill["count"].values}))
    R.h2("Signal decay: rank correlation by look-ahead (learn period)")
    ich = ic_by_horizon(m, F_ok, target, train_frac).reindex(rk_["col"])
    dec = pd.DataFrame({"Indicator": [META[c][0] for c in ich.index]})
    for col in ich.columns:
        dec[f"{col}m"] = ["-" if np.isnan(v) else f"{v * 100:+.0f}%" for v in ich[col].values]
    R.table(dec, font=7)
    R.note("Longer look-aheads have far fewer independent observations, so strong-looking numbers on the right are less trustworthy.")


def _rp_oos(R):
    if not (good or bad):
        R.box("No environment passed the strength filter, so the rules have no opinion. Lower the minimum strength (or raise the FDR limit) in the sidebar.", "warn")
    else:
        R.h2("Learn period (in-sample, flattering by construction)")
        R.table(fmt_summary(summarize(fwd, verdict, learn_idx, h, fdd), True))
        R.h2("Validation period (unseen by the rules)")
        R.table(fmt_summary(summarize(fwd, verdict, val_idx, h, fdd), True))
        if REVEAL:
            R.h2("Final test (never used for any choice, if this is the first reveal)")
            tsum = summarize(fwd, verdict, test_idx, h, fdd)
            R.table(fmt_summary(tsum, True))
            tb = tsum.set_index("Group")
        else:
            R.note("Final test hidden (research mode).")
            tb = summarize(fwd, verdict, val_idx, h, fdd).set_index("Group")
        g_ = _rp_gap_text(tb, grade_name)
        if g_:
            R.box(*g_)
    R.h2(f"Expanding-window rules (refitted every {RULE_STEP} months, fully out-of-sample)")
    if wfr.empty:
        R.p("Not enough history to refit the rules walk-forward.")
        return
    sw = summarize(fwd, wfr["v"], unseen_idx.intersection(wfr.index), h, fdd)
    R.table(fmt_summary(sw, True))
    g_ = _rp_gap_text(sw.set_index("Group"), f"Expanding-window rules, {unseen_lbl}")
    R.box(*g_) if g_ else R.note("Too few Favorable or Unfavorable months to compare.")


def _rp_regimes(R):
    if reg is None:
        R.p("Regime classification needs at least one growth indicator and one inflation indicator.")
    else:
        R.h2("Growth × inflation regimes")
        R.table(fmt_summary(summarize(fwd, reg, ev_r, h, fdd, order=REGIMES), True))
        R.p(f"**Today:** {reg.iloc[-1]}")
        tr_df, tr_cur, _ = regime_transitions(reg, fwd, ev_r)
        R.p(f"**Current transition:** {tr_cur}" if tr_cur else "No regime change in the last 3 months.")
        if not tr_df.empty:
            R.table(pd.DataFrame({"Transition": tr_df["Transition"], "Months": tr_df["Months"], f"Avg after {h}m": tr_df["Avg"].map(pc),
                                  "Median": tr_df["Median"].map(pc), "% positive": tr_df["Win"].map(lambda v: pc(v, False))}))
            R.note("Overlapping windows, so treat n as optimistic.")
    R.h2("10 most similar historical environments (descriptive, not tradable)")
    R.table(pd.DataFrame({"Date": ana["Date"], "Closeness": ana["Closeness"].map(lambda v: f"{v:.0f}%"), f"Outcome after {h}m": ana["After"].map(pc)}))
    R.h2(f"Predictive analogues (tradable signal, expanding database, {unseen_lbl})")
    if wa.empty:
        R.p("Not enough history for the predictive analogue test.")
        return
    ev_a = unseen_idx.intersection(wa.index)
    sa = summarize(fwd, ana_ver, ev_a, h, fdd)
    R.table(fmt_summary(sa, True))
    ic_a = rank_ic(wa["med"].reindex(ev_a), fwd.reindex(ev_a))
    g_ = _rp_gap_text(sa.set_index("Group"), f"Predictive analogues, {unseen_lbl}")
    R.box(g_[0] + f" Rank IC of the analogue median = {ic_a:+.2f}.", g_[1]) if g_ else R.note(f"Too few Favorable or Unfavorable analogue calls to compare (rank IC {pc(ic_a)}).")


def _rp_backtest(R, charts):
    if not all(k in globals() for k in ("pairs", "aux", "val_bt", "test_bt", "grade_bt")):
        R.p("The backtest was not run: the validation or final-test period is too short. Move the start date earlier or lower the learn / validation share.")
        return
    rf_ = cash_rate(m)
    full_ = {sname: bt_one(m, v_, aux, mode, bps, impl_key) for sname, v_ in pairs}
    first_ = pairs[0][0]
    R.p(f"Strategy: **{mode}**. The signal at each month-end sets the position for the next month; idle money earns the T-bill rate. Cost: {bps} bps per 100% traded. "
        f"Implementation: **{impl_key}**. Validation {val_bt[0]:%b %Y} to {val_bt[-1]:%b %Y} ({len(val_bt)} months)"
        + (f", final test {test_bt[0]:%b %Y} to {test_bt[-1]:%b %Y} ({len(test_bt)} months)." if REVEAL else ". Final test hidden (research mode)."))
    if impl_key == "Futures":
        R.box("Futures mode uses Yahoo's continuous front-month series, so roll gaps and roll yield are baked into the returns. Treat it as a rough check.", "warn")
    last_curves, last_idx = None, None
    for lbl, idx_ in [("Validation (used for choices)", val_bt)] + ([("Final test (untouched)", test_bt)] if REVEAL else []):
        rows_, curves_ = [], {}
        for sname, _ in pairs:
            ret, turn, W = full_[sname]
            r_, ntr, inv = seg(ret, turn, W, idx_)
            curves_[sname] = r_
            rows_.append((sname, perf(r_, rf_, inv, ntr)))
        ws, wg = float(full_[first_][2]["btc"].reindex(idx_).mean()), float(full_[first_][2]["gold"].reindex(idx_).mean())
        sm_ = static_mix(m, ws, wg, idx_, impl_key)
        curves_["No-timing mix (same avg exposure)"] = sm_
        rows_.append(("No-timing mix (same avg exposure)", perf(sm_, rf_, ws + wg, 1)))
        for bname, bret in bench_curves(m, idx_, impl_key).items():
            curves_[bname] = bret
            rows_.append((bname, perf(bret, rf_, 0.0 if bname.startswith("Cash") else 1.0, 0 if bname.startswith("Cash") else 1)))
        R.h2(lbl)
        R.table(perf_table(rows_))
        last_curves, last_idx, last_lbl = curves_, idx_, lbl
    R.note("The no-timing mix holds the same average bitcoin / gold / cash split as the strategy every month. If the strategy cannot beat it, the signal added nothing: "
           "the result came from being less invested, not from timing. A short test period and few trades make all numbers noisy.")
    if charts and last_curves:
        R.image(growth_chart(last_curves, last_idx, f"Growth of 1 unit: {last_lbl}"), h_px=480)

    pm = rank_by if find_best else "Sharpe"
    if run_perm:
        R.h2(f"Randomization test on the {grade_name.lower()} ({pm})")
        prow = []
        for sname, v_ in pairs:
            ret, turn, W = full_[sname]
            obs = perf(seg(ret, turn, W, grade_bt)[0], rf_)[pm]
            p_, med_, p95_ = perm_p(shift_null(m, v_, aux, mode, bps, impl_key, grade_bt.values, pm), obs)
            prow.append({"Signal": sname, f"Observed {pm}": fm(pm, obs), "Random median": fm(pm, med_), "Random 95th pct": fm(pm, p95_),
                         "p-value": "-" if pd.isna(p_) else f"{p_:.2f}", "Verdict": "beats luck" if (pd.notna(p_) and p_ <= 0.10) else "not distinguishable from luck"})
        R.table(pd.DataFrame(prow))
        R.note("p-value = share of time-shifted signals that scored at least as well (≤ 0.10 suggestive, ≤ 0.05 better). "
               + ("A fair test only for a strategy chosen before seeing the final test." if REVEAL else "In research mode this runs on the validation period, which you may have used for choices, so it is optimistic."))
    R.h2(f"Cost stress (first signal, {grade_name.lower()})")
    crow = []
    for b_ in sorted({0, bps, 50, 100}):
        rr, tt, WW = bt_one(m, pairs[0][1], aux, mode, b_, impl_key)
        pp = perf(seg(rr, tt, WW, grade_bt)[0], rf_)
        crow.append({"Cost (bps per 100% traded)": b_, "CAGR": pc(pp["CAGR"]), "Sharpe": fm("Sharpe", pp["Sharpe"]), "Max drawdown": pc(pp["Max drawdown"])})
    R.table(pd.DataFrame(crow))
    R.h2(f"Implementation check (first signal, {grade_name.lower()})")
    irow = []
    for ik in ("ETF", "Futures"):
        if ik == "ETF" and not (m["btc"].notna().sum() > 36 and m["gld"].notna().sum() > 36):
            continue
        rr, tt, WW = bt_one(m, pairs[0][1], aux, mode, bps, ik)
        pp = perf(seg(rr, tt, WW, grade_bt)[0], rf_)
        irow.append({"Implementation": ik, "CAGR": pc(pp["CAGR"]), "Sharpe": fm("Sharpe", pp["Sharpe"]), "Max drawdown": pc(pp["Max drawdown"])})
    R.table(pd.DataFrame(irow))
    if find_best and "lrows" in globals() and len(lrows):
        R.h2(f"Strategy leaderboard (ranked by {rank_by} on validation only, top 10)")
        R.table(pd.DataFrame(lrows).head(10))
        R.note("Many candidates were compared on the validation period, so the winner is flattered by selection. Trust a strategy only if it also beats bitcoin buy & hold "
               "on the final test and survives the randomization test.")


def _rp_robust(R):
    n_t_ = len(stats)
    if n_t_:
        thr_t = t_thr if cfg[0] == "t" else 1.5
        p_thr_ = math.erfc(thr_t / math.sqrt(2))
        R.h2("Multiple testing: how many findings could be luck?")
        R.kv([("Environments tested", n_t_), (f"Pass |t| ≥ {thr_t:.1f}", int((stats["t"].abs() >= thr_t).sum())),
              ("Expected to pass by pure luck", f"{n_t_ * p_thr_:.1f}"),
              ("Survive FDR q ≤ 0.10 / 0.25", f"{int((stats['q'] <= 0.10).sum())} / {int((stats['q'] <= 0.25).sum())}")])
        bq = stats.sort_values("q").head(8)
        R.table(pd.DataFrame({"Environment": [state_label(r.col, r.state) for r in bq.itertuples()], "Indicator": [META[r.col][0] for r in bq.itertuples()],
                              "HAC t": [f"{v:+.1f}" for v in bq["t"]], "Raw p": [f"{v:.3f}" for v in bq["p"]], "FDR q": [f"{v:.2f}" for v in bq["q"]],
                              "Both halves?": ["Yes" if v else "No" for v in bq["consistent"]]}))
        R.note("If the number passing is close to the number expected by luck, the rules are mostly noise. The luck estimate treats tests as independent, so it is rough.")
    if fac_mode and fac_members:
        R.h2("Factor composition")
        R.table(pd.DataFrame([{"Factor": META_ALL[k][0], "Pillar": META_ALL[k][4], "Members": ", ".join(META_ALL[c][0] for c in v)} for k, v in fac_members.items()]))
    if _rp_v("abl"):
        R.h2("Ablations: does each add-on help out of sample?")
        R.table(pd.DataFrame(abl))
        R.note("Diagnostics, not a menu: picking the best row would be selection on the grading period. Differences of a few points are within noise.")
    if _rp_v("stab"):
        R.h2("Is the factor structure stable?")
        R.table(pd.DataFrame(stab))
    if _rp_v("srows"):
        R.h2("Signal comparison (graded period)")
        R.table(pd.DataFrame(srows))
    sens_ = _rp_v("sens")
    if sens_ is not None:
        R.h2(f"Parameter sensitivity ({grade_name}: Favorable minus Unfavorable, pts)")
        tab = pd.DataFrame({"Learn share \\ min HAC t": [str(i) for i in sens_.index]})
        for c_ in sens_.columns:
            tab[str(c_)] = ["-" if pd.isna(v) else f"{v:+.1f}" for v in sens_[c_].values]
        R.table(tab)
        zs = sens_.values
        nz, npos = int(np.isfinite(zs).sum()), int((zs[np.isfinite(zs)] > 0).sum())
        if nz:
            R.box(f"{npos} of {nz} parameter settings keep Favorable ahead of Unfavorable on the {grade_name.lower()}. "
                  + ("A robust finding survives most settings." if npos / nz >= 0.8 else "If the result only works at one setting, it is fragile."), "ok" if npos / nz >= 0.8 else "warn")
    ph = _rp_v("pairs_hi")
    R.h2("Redundancy between indicators")
    if ph:
        R.p("Highly overlapping pairs (|rank correlation| ≥ 0.8). They vote twice for the same story:")
        R.bullets([f"{META[a][0]} ↔ {META[b][0]} ({r:+.2f})" for a, b, r in sorted(ph, key=lambda x: -abs(x[2]))])
    else:
        R.p("No pair of selected indicators has |rank correlation| ≥ 0.8.")


def _rp_explorer(R):
    ix = _rp_v("ind_x", next(iter(META)))
    sx = _rp_v("st_x", "High")
    allowed = S.index if REVEAL else S.index[S.index <= test_idx[0] - pd.DateOffset(months=12)]
    sel = S.index[(S[ix] == sx).fillna(False).values].intersection(allowed)
    R.p(f"Selection shown in the app when the report was built: **{META[ix][0]}**, state **{state_label(ix, sx)}**. "
        f"{len(sel)} months ({len(sel) / len(S):.0%} of history) were in this state. Now: **{state_label(ix, S.loc[now, ix])}** (value {fmt_val(ix, F_ok.loc[now, ix])}).")
    rows_ = []
    for h_ in (1, 3, 6, 12):
        row = {"Horizon": f"{h_} months"}
        for nm_, kd in (("Bitcoin", "ret"), ("Gold", "goldabs"), ("Bitcoin − gold", "gold")):
            tr_ = target_returns(m, F_ok.index, h_, kd)
            r, base_r = tr_.reindex(sel).dropna(), tr_.reindex(allowed).dropna()
            row[nm_] = "-" if r.empty else f"{r.mean():+.1%} ({(r > 0).mean():.0%} up, n={len(r)}; all months {base_r.mean():+.1%})"
        dd_ = fwd_drawdown(m["btc"], h_).reindex(sel).dropna()
        row["Bitcoin typical max drawdown"] = "-" if dd_.empty else f"{dd_.median():.0%}"
        rows_.append(row)
    R.table(pd.DataFrame(rows_))
    R.note("The Explorer is interactive in the app: change the indicator and state there, then rebuild the report to capture a different selection.")


def _rp_expectations(R):
    xc = _rp_v("_XC", [c for c in EXP_KEYS if c in F_full.columns])
    if not xc:
        R.p("No expectation indicator could be built (the Treasury / FRED series did not download).")
        return
    icon_txt = {"Market-implied": "Market-implied", "Published model": "Published model", "Model-implied": "Model-implied"}
    last_ = F_full[xc].dropna(how="all").index[-1]
    R.p(f"Eleven forward-looking indicators, all point-in-time. **Latest read: {last_:%b %Y}.** Market-implied = computed from the Treasury curve; Published model = Cleveland Fed; "
        "Model-implied = walk-forward ridge estimate (no free point-in-time survey history exists), so it adds convenience, not new information.")
    rows_ = []
    for c in xc:
        s_ = F_full[c].dropna()
        if s_.empty:
            rows_.append({"Indicator": META_ALL[c][0], "Type": icon_txt[EXP_INFO[c]["method"]], "Latest": "not available", "Change vs 3 months earlier": "-", "Textbook sign for bitcoin": _SIGN_TXT[EXP_INFO[c]["btc"]]})
            continue
        now_v, prev = s_.iloc[-1], F_full[c].shift(3).loc[s_.index[-1]]
        rows_.append({"Indicator": META_ALL[c][0], "Type": icon_txt[EXP_INFO[c]["method"]], "Latest": _fmt(c, now_v),
                      "Change vs 3 months earlier": "-" if pd.isna(prev) else _fmt(c, now_v - prev, True), "Textbook sign for bitcoin": _SIGN_TXT[EXP_INFO[c]["btc"]]})
    R.table(pd.DataFrame(rows_))

    def val(c):
        s_ = F_full[c].dropna() if c in F_full else pd.Series(dtype=float)
        return s_.iloc[-1] if len(s_) else np.nan

    lines = []
    d_ff = val("ex_ff_chg")
    if pd.notna(d_ff):
        lines.append(f"**Fed:** the curve prices {'cuts' if d_ff <= -0.25 else ('hikes' if d_ff >= 0.25 else 'roughly no change')} ({d_ff:+.2f} pts over 12 months, to about {val('ex_ff12'):.2f}%).")
    if pd.notna(val("ex_real")):
        lines.append(f"**Real policy rate in 12 months:** about {val('ex_real'):.2f}% (expected inflation {val('ex_infl'):.2f}%).")
    rp = val("ex_recess")
    if pd.notna(rp):
        lines.append(f"**Recession risk (yield-curve model):** {rp:.0f}% ({'low' if rp < 20 else ('moderate' if rp < 40 else 'elevated')}).")
    if pd.notna(val("ex_growth")) and pd.notna(val("ex_unrate")):
        lines.append(f"**Economy (model-implied):** growth about {val('ex_growth'):.1f}%, unemployment about {val('ex_unrate'):.1f}% in 12 months.")
    u_ = val("ex_usd")
    if pd.notna(u_):
        lines.append(f"**Dollar (model-implied):** {'stronger' if u_ >= 0.01 else ('weaker' if u_ <= -0.01 else 'about flat')} ({u_:+.1%} over 12 months; low-confidence kind of forecast).")
    if lines:
        R.bullets(lines)
    sup = uns = neu = mix = 0
    for c in xc:
        sg = EXP_INFO[c]["btc"]
        p_ = P_full[c].dropna().iloc[-1] if c in P_full and P_full[c].notna().any() else np.nan
        if sg == 0:
            mix += 1
        elif pd.isna(p_):
            continue
        elif (p_ >= 2 / 3 and sg > 0) or (p_ <= 1 / 3 and sg < 0):
            sup += 1
        elif (p_ >= 2 / 3 and sg < 0) or (p_ <= 1 / 3 and sg > 0):
            uns += 1
        else:
            neu += 1
    R.p(f"**Textbook read for bitcoin:** {sup} expectation(s) point in bitcoin's favor, {uns} against, {neu} in the middle of their range, {mix} have no clear textbook sign. "
        "This is a checklist, not evidence; the table below shows what the data in this sample says.")

    R.h2(f"Do the expectations line up with bitcoin's {h}-month outcome? ({target_name.lower()})")
    rows_ = []
    for c in xc:
        x_ = F_full[c].reindex(fwd.index)
        xl, yl, xt, yt = x_.loc[learn_idx], fwd.loc[learn_idx], x_.loc[unseen_idx], fwd.loc[unseen_idx]
        ic_l, ic_t = rank_ic(xl, yl), rank_ic(xt, yt)
        dl = pd.concat([xl, yl], axis=1).dropna()
        t_l = nw_t(dl.iloc[:, 1].rank().values, dl.iloc[:, 0].rank().values, h) if len(dl) > 8 else np.nan
        same = bool(pd.notna(ic_l) and pd.notna(ic_t) and np.sign(ic_l) == np.sign(ic_t))
        status = "-" if (pd.isna(ic_l) or pd.isna(ic_t)) else ("Flips on unseen data" if not same else ("Consistent and significant" if abs(t_l) >= 2 else "Consistent but weak"))
        sg = EXP_INFO[c]["btc"]
        rows_.append({"Indicator": META_ALL[c][0], "Learn months": len(dl), "Rank IC (learn)": "-" if pd.isna(ic_l) else f"{ic_l:+.2f}",
                      f"Rank IC ({unseen_lbl})": "-" if pd.isna(ic_t) else f"{ic_t:+.2f}", "HAC t (learn)": "-" if pd.isna(t_l) else f"{t_l:+.1f}", "Status": status,
                      "Textbook sign": _SIGN_TXT[sg], "Matches textbook?": "n/a" if sg == 0 or pd.isna(ic_l) else ("Yes" if np.sign(ic_l) == sg else "No")})
    R.table(pd.DataFrame(rows_))
    R.note("Eleven indicators are tested, so a lone 'consistent' result can be luck. Expectations formed from the same curve are strongly correlated with each other and with the existing curve and real-yield indicators.")

    R.h2("Are the expectations any good? (what actually happened 12 months later)")
    gr = grade_expectations(m, F_full[xc], ev_r)
    if gr.empty:
        R.p("Not enough overlapping months to grade the expectations.")
    else:
        nm = {"ex_ff12": "Fed funds, 12M ahead", "ex_y10": "10Y yield, 12M ahead", "ex_growth": "GDP growth", "ex_unrate": "Unemployment rate",
              "ex_usd": "Dollar 12M change", "ex_m2": "M2 growth", "ex_liq": "Fed net-liquidity growth"}
        R.table(pd.DataFrame({"Expectation": [f"{nm[k]} ({EXP_INFO[k]['method']})" for k in gr["k"]], "Months graded": gr["n"],
                              "Avg error, expectation": gr["mae_m"].map(lambda v: f"{v:.3f}"), "Avg error, 'nothing changes'": gr["mae_n"].map(lambda v: f"{v:.3f}"),
                              "Skill vs 'nothing changes'": gr["skill"].map(lambda v: pc(v)), "IC, expected vs realised change": gr["ic"].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}")}))
        beat = int((gr["skill"] > 0).sum())
        R.box(f"{beat} of {len(gr)} expectations had a smaller average error than 'nothing changes'. Windows overlap, so independent observations are roughly months ÷ 12. "
              + ("Where skill is negative, treat the indicator as a description of the model's view, not as a forecast." if beat < len(gr) else ""),
              "ok" if beat >= len(gr) * 0.6 else "warn")
    R.h2("Definitions and sources")
    R.table(pd.DataFrame([{"Indicator": META_ALL[c][0], "Type": EXP_INFO[c]["method"], "How it is built": EXP_INFO[c]["how"], "Source": EXP_INFO[c]["src"],
                           "First month": F_full[c].first_valid_index().strftime("%b %Y") if F_full[c].notna().any() else "-"} for c in xc]), font=7, widths=[3, 2, 7.2, 3.2, 2])


def _rp_projection(R, charts):
    P_ = PROJ_REPORT
    if not P_.get("head"):
        R.p("The projection was not produced (not enough history with the current settings, or the projection data did not download).")
        return
    hd, L_, S0_ = P_["head"], P_["L"], P_["S0"]
    R.p(f"Machine-learning fan of possible bitcoin price routes for the next **{L_} months** (as of {P_['last']:%b %Y}). Model: **{P_['model']}**; features: **{P_['feats']}**; "
        f"trust in the model: **{P_['trust']}%** (0% = history only). Each month's spread is the predicted distribution of bitcoin's return (10 / 25 / 50 / 75 / 90% levels).")
    R.kv([(f"Median price in {L_} months", f"${hd['median']:,.1f} ({hd['chg']:+.1%} vs ${S0_:,.1f} now)"), (f"80% range in {L_} months", f"${hd['lo']:,.0f} to ${hd['hi']:,.0f}"),
          (f"Chance higher in {L_} months", f"{hd['up']:.0%}"),
          (f"Chance of a {PROJ_DD:.0%}+ drawdown on the way", f"{hd['dd']:.0%} (typical worst fall {hd['dd_med']:.0%})")])
    if charts and P_.get("fig") is not None:
        R.image(P_["fig"], title=f"Bitcoin route fan: {P_['model']}, trust {P_['trust']}%")
    if P_.get("table") is not None:
        R.h2("Projection by horizon")
        R.table(P_["table"])
    if P_.get("score") is not None:
        R.h2(f"Does the model beat plain history? (walk-forward, {P_['unseen_lbl']} months only)")
        R.table(P_["score"])
        ok_ = P_["share"] >= 0.7 and P_["avg_skill"] > 0.02
        R.box((f"The model beat history-only on {P_['n_good']} of {P_['n_hor']} horizons (average skill {P_['avg_skill']:+.1%}). Some real information, but the test is short and windows overlap: read the median tilt as a mild lean, not a target."
               if ok_ else f"The model did NOT reliably beat history-only ({P_['n_good']} of {P_['n_hor']} horizons better, average skill {P_['avg_skill']:+.1%}). "
               "Treat the fan as a risk range and ignore any tilt in the median."), "ok" if ok_ else "warn")
        if P_["cov"] < 0.70 or P_["cov"] > 0.90:
            R.box(f"The 80% range contained the outcome {P_['cov']:.0%} of the time on unseen months, so the fan is {'too narrow' if P_['cov'] < 0.70 else 'too wide'}.", "warn")
        R.note("Skill = reduction in pinball loss versus the history-only fan (positive = better). Coverage should be near 80% / 50%. Direction hit: 50% = coin flip. "
               "Independent observations are roughly months ÷ horizon.")
    if P_.get("trust_tab") is not None:
        R.h3("Skill by trust level (a diagnostic, not a menu)")
        R.table(P_["trust_tab"])
    if P_.get("compare") is not None:
        R.h2("Do the expectation indicators help the projection?")
        R.table(P_["compare"])
    elif _rp_v("_XC"):
        R.note("The with / without expectations comparison was not run in the app (tick the box in the Projection tab, then rebuild the report to include it).")
    if P_.get("drivers") is not None:
        R.h2(f"What is pushing the {P_['kd']}-month projection today?")
        dv = P_["drivers"]
        R.table(pd.DataFrame({"Feature": [str(i) for i in dv.index], "Contribution (pct. points of expected log return)": [f"{v:+.2f}" for v in dv.values]}))
        R.note("Correlated indicators share credit unpredictably, so read this as a story about the inputs, not as causes.")
    R.box("Not a price target and not financial advice. The fan describes how wide bitcoin's routes have been after similar conditions, adjusted by a model that may have no skill.", "warn")


def _rp_ml(R):
    if not run_ml:
        R.p("Machine learning was switched off in the sidebar.")
        return
    R.p(f"Each model estimates the probability that {tgt_txt} in {h} months, retrained every {step} months on data whose outcome was already known. "
        f"Scores use {unseen_lbl} months. AUC 0.50 = coin flip, 0.55 to 0.60 = modest skill; the bracket is a 95% block-bootstrap range. "
        "Brier skill > 0 means the probabilities beat quoting the normal odds.")
    rows_ = []
    for name in ml:
        s_ = _rp_ml_stats(name)
        if s_:
            rows_.append({"Model": name, "AUC": f"{s_['auc']:.2f} [{s_['lo']:.2f} to {s_['hi']:.2f}]" if pd.notna(s_["lo"]) else f"{s_['auc']:.2f}",
                          "Brier skill": f"{s_['bss']:+.3f}", "Log loss": f"{s_['ll']:.3f}", "Rank IC": f"{s_['ic']:+.2f}", "Accuracy": f"{s_['acc']:.0%}",
                          "Naive": f"{s_['naive']:.0%}", "Favorable minus Unfavorable": "-" if np.isnan(s_["spr"]) else f"{s_['spr']:+.1f} pts"})
    R.h2(f"Out-of-sample comparison ({unseen_lbl} months)")
    R.table(pd.DataFrame(rows_))
    R.note("If logistic regression scores about the same as the forests, the extra complexity is not buying anything.")
    s_ = _rp_ml_stats(ml_choice)
    if not s_:
        R.p("Not enough history to train. Move the start date earlier.")
        return
    skill = pd.notna(s_["lo"]) and s_["lo"] > 0.5 and s_["bss"] > 0
    R.box(f"{ml_choice}: AUC {s_['auc']:.2f}" + (f" (95% range {s_['lo']:.2f} to {s_['hi']:.2f})" if pd.notna(s_["lo"]) else "") + f", Brier skill {s_['bss']:+.3f}. "
          + ("The AUC range sits above 0.50, so there is some skill, but the test period is short." if skill else
             "The AUC range includes 0.50 or the probabilities do not beat the normal odds: little reliable skill on unseen data. Don't rely on it."), "ok" if skill else "warn")
    R.h2(f"{ml_choice}: outcome after each call")
    R.table(fmt_summary(summarize(fwd, ml_verdict(s_["wk"], margin), s_["wk"].index, h, fdd), True))
    R.p(f"**Today's read:** {ml_p:.0%} chance that {tgt_txt} in {h} months, versus {ml_base:.0%} normal odds.")
    ser = _rp_v("ml_contrib")
    ttl = "Why today's prediction? (log-odds contribution of each indicator)" if ser is not None else "What the model uses most"
    ser = ml_contrib if ser is not None else _rp_v("ml_imp")
    if ser is not None:
        ser = ser.reindex(ser.abs().sort_values(ascending=False).index).head(12)
        R.h2(ttl)
        R.table(pd.DataFrame({"Indicator": [META[c][0] if c in META else str(c) for c in ser.index], "Value": [f"{v:+.3f}" for v in ser.values]}))
        R.note("Positive pushes toward the target being positive, negative pushes away. Importance is not causality.")


def _rp_scan(R):
    sc_ = _rp_v("scan")
    if sc_ is None:
        R.p("The look-ahead scan was not run (tick 'Auto-find best look-ahead' in the sidebar).")
        return
    R.p(f"Each look-ahead is fitted on the first 70% of the learning period and scored on its last 30%. The selected look-ahead is **{h} months**, then locked. "
        "Validation and final test are never used here. Stable IC is the worse of the two validation halves.")
    d_ = sc_.copy()
    out = pd.DataFrame({"Look-ahead": d_["h"].map(lambda x: f"{x} months" + (" (selected)" if x == h else ""))})
    for c, nm in (("rules_ic", "Rules IC"), ("ml_ic", "ML IC"), ("combined", "Average IC"), ("stable", "Stable IC")):
        out[nm] = d_[c].map(lambda v: "-" if pd.isna(v) else f"{v:+.2f}")
    out["Independent obs."] = (sc_["n_val"] / sc_["h"]).round(0).astype(int)
    R.table(out)
    R.note("Longer look-aheads have far fewer independent observations. If every IC is near zero, there is no reliable horizon, and that is a valid answer.")


def _rp_now(R):
    R.p(f"Each market gets its own rules and ML read for the next {h} months, using the same indicators and settings. Favorable / Unfavorable describes how similar environments "
        "played out historically; it is not a buy or sell instruction. These reads always use plain return targets.")
    for k in NOW_KINDS:
        r = NOW[k]
        R.h2(f"{NOW_NAMES[k]}: {REL_TXT[r['label']] if k == 'gold' else r['label']}")
        it = [f"**Confidence: {r['conf']}** ({r['agree']} of {r['n_votes']} evidence channels agree): " + ", ".join(f"{n_} {_RP_VOTE[v]}" for n_, v in r["votes"].items()) + ".",
              "Macro models: " + ", ".join(f"{n_} {_RP_VOTE[v]}" for n_, v in r["comps"].items()) + ".", f"Rules: {r['rules']} (score {r['score']:+.1f})."]
        if r["ml"]:
            it.append(f"ML: {r['ml']['p']:.0%} chance {NOW_PHRASE[k]} in {h}m (normal {r['ml']['base']:.0%}) → {r['ml']['verdict']}.")
        it.append(f"Similar past periods (descriptive): median {pc(r['ana_med'])}, {r['ana_pos']:.0%} positive.")
        it.append(f"**Evidence on unseen data: {r['ev_level']}.** {r['ev_text']}")
        if r["ml"] and pd.notna(r["ml"]["auc"]):
            it.append(f"ML on unseen data: AUC {r['ml']['auc']:.2f}, Brier skill {r['ml']['bss']:+.3f}.")
        if r["why_good"]:
            it.append("**Helping now:** " + "; ".join(r["why_good"][:4]))
        if r["why_bad"]:
            it.append("**Hurting now:** " + "; ".join(r["why_bad"][:4]))
        R.bullets(it)
    R.note("Confidence combines whether the two evidence channels (macro models and price trend) point the same way and whether the rules beat chance on unseen data. It is a rough guide, not a probability.")
    R.box("Educational only, not financial advice. This reads macro conditions only. It does not know about valuation, news, taxes, your goals or your time horizon, "
          "and relationships that held in the past can stop working.", "warn")


def build_word_report(charts=True):
    """Returns (docx bytes, list of section names that could not be built)."""
    _RP_IMG_OK[0] = bool(charts)
    R, fails = _Rep(), []
    t = R.d.add_paragraph()
    R._runs(t, "Bitcoin Macro Environment Report", size=26, color="1F3864", bold=True)
    R.p(f"Results of all tabs · data through {m.index[-1]:%B %Y} · prepared {pd.Timestamp.today():%d %B %Y}", size=11, color="595959")
    R.p(f"Predicting: {target_name}, {h} months ahead · mode: {'locked evaluation, final test revealed' if REVEAL else 'research (final test hidden)'}", size=10, color="595959")
    plan = [("Executive summary", _rp_summary), ("Settings and data", _rp_settings), ("Dashboard", _rp_dash), ("Guide: how this works", _rp_guide),
            ("Environments", _rp_env), ("Indicator ranking", _rp_rank), ("Out-of-sample", _rp_oos), ("Regimes & analogues", _rp_regimes),
            ("Backtest", lambda R_: _rp_backtest(R_, charts)), ("Robustness", _rp_robust), ("Explorer", _rp_explorer), ("Expectations", _rp_expectations),
            ("Projection", lambda R_: _rp_projection(R_, charts)), ("Machine learning", _rp_ml), ("Look-ahead scan", _rp_scan), ("Bitcoin & gold now", _rp_now)]
    for i, (name, fn) in enumerate(plan, 1):
        R.h1(f"{i}. {name}")
        if i == 1:
            R.d.paragraphs[-1].paragraph_format.page_break_before = False
        try:
            fn(R)
        except Exception as e:  # noqa: BLE001  (one broken section must not lose the whole report)
            fails.append(f"{name} ({type(e).__name__}: {e})")
            R.box(f"This section could not be built ({type(e).__name__}: {e}). The other sections are unaffected.", "bad")
    return R.save(), fails


# ------------------------------------------------------------------ report tab
with T["report"]:
    st.subheader("📄 Word report: the results of all tabs in one document")
    st.caption("Builds a Word (.docx) file with an executive summary, the settings used, and the results of every tab (Dashboard to Bitcoin & gold now): tables, verdicts and the main charts. "
               "It reports exactly what the app shows for the current settings, including the research / locked mode: in research mode the final test stays hidden. "
               "Press the button after you have finished changing settings; the report is not rebuilt on every click.")
    if not _HAS_DOCX:
        st.error("The python-docx package is not installed. Add `python-docx` to requirements.txt (and optionally `kaleido` for charts), then restart the app.")
    else:
        want_charts = st.checkbox("Include charts (needs the optional `kaleido` package; without it the report is built without charts)", True, key="rep_charts")
        sig = repr((str(m.index[-1]), start, target, h, REVEAL, sorted(META), fac_mode, pit, cfg, use_exp, run_ml, ml_choice, mode, impl_key, bps,
                    proj_model, proj_feats, proj_len, proj_trust, want_charts))
        if st.button("📝 Build the Word report", key="rep_build"):
            with st.spinner("Collecting results from every tab and writing the Word file..."):
                _data, _fails = build_word_report(want_charts)
            st.session_state["rep_saved"] = (sig, _data, _fails, pd.Timestamp.now())
        saved = st.session_state.get("rep_saved")
        if saved:
            sig0, data_, fails_, ts_ = saved
            if sig0 != sig:
                st.warning("Settings or data have changed since this report was built. Rebuild it to match what the tabs show now.")
            st.success(f"Report ready ({len(data_) / 1024:,.0f} KB, built {ts_:%H:%M:%S}).")
            st.download_button("⬇️ Download the report (.docx)", data_, file_name=f"bitcoin_report_{now:%Y_%m}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="rep_dl")
            if fails_:
                st.warning("Some sections could not be built and contain an error note instead: " + "; ".join(fails_))
            if want_charts and not _RP_IMG_OK[0]:
                st.info("Charts were left out because the `kaleido` package is not installed. Add `kaleido` to requirements.txt to include them.")
